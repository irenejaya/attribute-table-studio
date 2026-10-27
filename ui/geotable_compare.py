# -*- coding: utf-8 -*-
# Attribute Table Studio - https://github.com/irenejaya/attribute-table-studio
# Copyright (C) 2026 Irene Jaya. Licensed under the GNU GPL v3 with an
# attribution term (section 7(b)): see LICENSE and ADDITIONAL_TERMS.md.
"""
Geometry and Table Compare - comparison panel (Attribute Table Studio tool 3)

Compares an OLD (baseline) layer against a NEW (updated) layer:
- Qt5 / Qt6 compatible (QGIS 3.40+ and QGIS 4.x) with light / dark theme support
- Pre-flight validation (duplicate IDs, field structure, null checks)
- Optional geometry property comparison (length / area / centroid / vertices)
- Colour-coded results with status filters, live counts and text search
- Review workflow (accept / reject changes) and Excel / CSV export

The comparison logic is unchanged from the standalone plugin (v1.1); the
interface uses the shared Attribute Table Studio theme, buttons and status strip
and is shown in a window or dock by ui/tool_host.py.
"""

import os
import csv
import math
from collections import Counter

from qgis.PyQt.QtCore import (
    Qt, QDate, QDateTime, QRect, QTimer, QSize, QStringListModel,
)
from qgis.PyQt.QtGui import QColor, QBrush, QFont, QFontMetrics, QPalette, QPen
from qgis.PyQt.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QGridLayout, QComboBox, QLabel,
    QTableWidget, QTableWidgetItem, QHeaderView, QCheckBox, QFileDialog,
    QMessageBox, QAbstractItemView, QWidget, QFrame, QApplication, QLineEdit,
    QStackedWidget, QSizePolicy, QScrollArea, QCompleter, QMenu,
    QDoubleSpinBox, QProgressBar, QPushButton, QStyle, QStyledItemDelegate,
    QStyleOptionViewItem,
)
from qgis.core import (
    Qgis, QgsWkbTypes, NULL, QgsGeometry, QgsFeature, QgsCsException,
    QgsVectorLayer, QgsProject, QgsRectangle, QgsCoordinateTransform,
    QgsSymbol, QgsSingleSymbolRenderer, QgsFeatureRequest, QgsSpatialIndex,
    QgsVectorFileWriter, QgsCategorizedSymbolRenderer, QgsRendererCategory,
    QgsMarkerSymbol,
)

from ..core import fields as FK
from ..core.compat import SELECT_SET
from . import icons
from . import widgets as W
from .theme import Metrics, Theme, stylesheet

try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        Font as XlFont, PatternFill, Alignment, Border, Side,
    )
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# ----------------------------------------------------------------------------
# Enums (QGIS 3.40+ and 4.x both have the scoped Qgis enums)
# ----------------------------------------------------------------------------
_VECTOR_FILTER = Qgis.LayerFilter.VectorLayer
_GEOM_POINT = Qgis.GeometryType.Point
_GEOM_LINE = Qgis.GeometryType.Line
_GEOM_POLYGON = Qgis.GeometryType.Polygon

_ROLE_USER = Qt.ItemDataRole.UserRole
_ROLE_KIND = Qt.ItemDataRole.UserRole + 1     # "changed" for pink cells
_ROLE_DECISION = Qt.ItemDataRole.UserRole + 2  # a changed value's decision
_ROLE_TOOLTIP = Qt.ItemDataRole.ToolTipRole

STATUSES = ("Added", "Deleted", "Modified", "Unchanged")
REVIEWABLE = ("Added", "Modified")
AUTO_SKIP_FIELDS = {
    'fid', 'id', 'objectid', 'gid',
    'created_date', 'modified_date', 'timestamp'
}
COL_STATUS = 0
COL_DECISION = 1
COL_FIRST_FIELD = 2
MAX_SUGGESTIONS = 5000
DEFAULT_GEOM_TOLERANCE = 0.001
PLUGIN_NAME = "Geometry and Table Compare"
MAP_GROUP_NAME = PLUGIN_NAME
REVIEW_LAYER_NAME = f"{PLUGIN_NAME} – Review points"

MODE_TIP = (
    "How multipart geometry is compared.\n\n"
    "Singlepart: every feature is split into its parts and each part is\n"
    "matched to the same part in the other layer. Review points mark each\n"
    "part that was added, removed or modified.\n\n"
    "Multipart: each feature is compared as one whole shape. Review points\n"
    "mark every connected piece that changed."
)

# Review-point attributes. Every name is <= 10 characters, so ESRI Shapefile
# (.dbf) never truncates them and the .shp and .gpkg outputs are identical.
REVIEW_FIELDS = (
    ("join_id", "string(80)"),    # value of the Join field
    ("change", "string(10)"),     # added / removed / modified
    ("level", "string(10)"),      # feature / part / region
    ("old_part", "integer"),      # part number in OLD (modified / removed)
    ("new_part", "integer"),      # part number in NEW (modified / added)
    ("chg_size", "double"),       # size of the change (see chg_unit)
    ("chg_unit", "string(10)"),   # area / length / points
)


# ============================================================================
# GEOMETRY DIFFERENCE (where did a - possibly huge multipart - shape change?)
# ============================================================================

def _geom_present(g):
    return g is not None and not g.isNull() and not g.isEmpty()


def _geom_unit(gtype):
    if gtype == _GEOM_POLYGON:
        return "area"
    if gtype == _GEOM_LINE:
        return "length"
    return "points"


def _geom_measure(g, gtype):
    """Area (polygon), length (line) or vertex count (point)."""
    if not _geom_present(g):
        return 0.0
    if gtype == _GEOM_POLYGON:
        return g.area()
    if gtype == _GEOM_LINE:
        return g.length()
    return float(sum(1 for _ in g.vertices()))


def _geom_parts(g):
    return g.constGet().partCount() if _geom_present(g) else 0


def _safe_difference(a, b):
    result = a.difference(b)
    if result is None or result.isNull():
        # invalid input geometry: repair and retry
        result = a.makeValid().difference(b.makeValid())
    return result


def _keep_type(g, gtype):
    """Parts of *g* with the layer's geometry type (drops slivers of a
    lower dimension), merged into a single multi-geometry."""
    if not _geom_present(g):
        return None
    parts = [p for p in g.asGeometryCollection()
             if _geom_present(p) and p.type() == gtype]
    if not parts:
        return None
    merged = QgsGeometry.collectGeometry(parts)
    return merged if _geom_present(merged) else None


def _topo_equal(a, b):
    """Topological equality (isGeosEqual is deprecated from QGIS 4.2)."""
    fn = getattr(a, "isTopologicallyEqual", None)
    if fn is not None:
        try:
            return bool(fn(b))
        except (TypeError, ValueError):
            pass                         # fall back to isGeosEqual
    return bool(a.isGeosEqual(b))


def geometry_diff(old_g, new_g, tolerance, gtype):
    """Return what was added to / removed from a geometry.

    Parts of NEW farther than *tolerance* from OLD are 'added'; parts of OLD
    farther than *tolerance* from NEW are 'removed'.
    """
    old_ok, new_ok = _geom_present(old_g), _geom_present(new_g)
    added = removed = None
    if old_ok and new_ok:
        if not _topo_equal(old_g, new_g):
            if tolerance > 0:
                added = _safe_difference(new_g, old_g.buffer(tolerance, 4))
                removed = _safe_difference(old_g, new_g.buffer(tolerance, 4))
            else:
                added = _safe_difference(new_g, old_g)
                removed = _safe_difference(old_g, new_g)
            added = _keep_type(added, gtype)
            removed = _keep_type(removed, gtype)
    elif new_ok:
        added = QgsGeometry(new_g)
    elif old_ok:
        removed = QgsGeometry(old_g)

    old_m = _geom_measure(old_g, gtype)
    added_m = _geom_measure(added, gtype)
    removed_m = _geom_measure(removed, gtype)
    return {
        "changed": bool(old_ok and new_ok and (added or removed)),
        "added": added,
        "removed": removed,
        "added_m": added_m,
        "removed_m": removed_m,
        "pct": ((added_m + removed_m) / old_m * 100.0)
        if old_ok and new_ok and old_m else None,
        "parts_old": _geom_parts(old_g),
        "parts_new": _geom_parts(new_g),
        "points": [],
        "parts_changed": 0,
    }


# ----------------------------------------------------------------------------
# Review points + part-by-part analysis
# ----------------------------------------------------------------------------

MAX_REGION_POINTS = 2000        # per feature, in 'whole feature' mode
MAX_PART_SCAN = 20000           # features scanned to detect multipart input


def _pieces(g, gtype):
    """Single-part pieces of *g* that have the layer's geometry type."""
    if not _geom_present(g):
        return []
    return [p for p in g.asGeometryCollection()
            if _geom_present(p) and p.type() == gtype]


def _review_point(g, gtype):
    """A point that lies on / in *g* (largest piece), for the user to visit."""
    pieces = _pieces(g, gtype) or ([g] if _geom_present(g) else [])
    if not pieces:
        return None
    if gtype == _GEOM_POLYGON:
        best = max(pieces, key=lambda p: p.area())
        pt = best.pointOnSurface()
    elif gtype == _GEOM_LINE:
        best = max(pieces, key=lambda p: p.length())
        pt = best.interpolate(best.length() / 2.0)
    else:
        pt = QgsGeometry(pieces[0])
    if not _geom_present(pt):
        pt = pieces[0].centroid()
    return pt if _geom_present(pt) else None


def _point_record(geom, gtype, change, level, part_old=None, part_new=None):
    pt = _review_point(geom, gtype)
    if pt is None:
        return None
    return {"geom": pt, "change": change, "level": level,
            "part_old": part_old, "part_new": part_new,
            "measure": _geom_measure(geom, gtype)}


def _feature_points(gd, gtype):
    """One point for a feature that exists in only one of the layers."""
    kind = "added" if gd["added"] else "removed"
    rec = _point_record(gd[kind], gtype, kind, "feature")
    return [rec] if rec else []


def _region_points(gd, gtype):
    """'Whole feature' mode: one point per connected added / removed piece."""
    out = []
    for kind in ("added", "removed"):
        for piece in _pieces(gd[kind], gtype)[:MAX_REGION_POINTS]:
            rec = _point_record(piece, gtype, kind, "region")
            if rec:
                out.append(rec)
    return out


def _parts_equal(a, b, tol):
    if _topo_equal(a, b):
        return True
    if tol <= 0:
        return False
    ra, rb = a.boundingBox(), b.boundingBox()
    if (abs(ra.xMinimum() - rb.xMinimum()) > tol
            or abs(ra.xMaximum() - rb.xMaximum()) > tol
            or abs(ra.yMinimum() - rb.yMinimum()) > tol
            or abs(ra.yMaximum() - rb.yMaximum()) > tol):
        return False
    h = a.hausdorffDistance(b)
    return 0 <= h <= tol


def _part_score(a, b, gtype):
    """How strongly OLD part *a* and NEW part *b* are the 'same' object."""
    try:
        if gtype == _GEOM_POLYGON:
            inter = a.intersection(b)
            ia = inter.area() if _geom_present(inter) else 0.0
            union = a.area() + b.area() - ia
            return ia / union if union > 0 else 0.0
        if gtype == _GEOM_LINE and a.intersects(b):
            h = a.hausdorffDistance(b)
            return 1.0 / (1.0 + max(h, 0.0))
    except (TypeError, ValueError, ZeroDivisionError):
        pass                             # unusable geometry: no match
    return 0.0


def _search_rect(g, tol):
    rect = QgsRectangle(g.boundingBox())
    rect.grow(max(tol, 0.0) + 1e-9)
    return rect


def _match_parts(old_parts, new_parts, tol, gtype):
    """Pair OLD and NEW parts of one feature.

    Returns (same, modified, only_old, only_new): lists of index pairs /
    indices.  Unchanged parts are matched first (so reordering is harmless);
    the rest are paired by overlap (polygons) or proximity (lines).
    """
    index = QgsSpatialIndex()
    add = getattr(index, "addFeature", None) or index.insertFeature
    for j, g in enumerate(new_parts):
        f = QgsFeature(j)
        f.setGeometry(g)
        add(f)

    used_new, same = set(), []
    unmatched_old = []
    for i, a in enumerate(old_parts):
        hit = None
        for j in index.intersects(_search_rect(a, tol)):
            if j not in used_new and _parts_equal(a, new_parts[j], tol):
                hit = j
                break
        if hit is None:
            unmatched_old.append(i)
        else:
            used_new.add(hit)
            same.append((i, hit))

    pairs = []
    for i in unmatched_old:
        a = old_parts[i]
        for j in index.intersects(_search_rect(a, tol)):
            if j in used_new:
                continue
            s = _part_score(a, new_parts[j], gtype)
            if s > 0:
                pairs.append((s, i, j))
    pairs.sort(reverse=True)

    modified, done_old = [], set()
    for _s, i, j in pairs:
        if i in done_old or j in used_new:
            continue
        done_old.add(i)
        used_new.add(j)
        modified.append((i, j))

    only_old = [i for i in unmatched_old if i not in done_old]
    only_new = [j for j in range(len(new_parts)) if j not in used_new]
    return same, modified, only_old, only_new


def geometry_diff_parts(old_g, new_g, tol, gtype):
    """Part-by-part analysis: every part of OLD is matched to a part of NEW.

    Gives the same result shape as :func:`geometry_diff` plus a review point
    for every part that was added, removed or modified.
    """
    old_parts = _pieces(old_g, gtype)
    new_parts = _pieces(new_g, gtype)
    same, modified, only_old, only_new = _match_parts(
        old_parts, new_parts, tol, gtype)

    added_geoms, removed_geoms, points = [], [], []
    for j in only_new:
        added_geoms.append(new_parts[j])
        rec = _point_record(new_parts[j], gtype, "added", "part",
                            part_new=j + 1)
        if rec:
            points.append(rec)
    for i in only_old:
        removed_geoms.append(old_parts[i])
        rec = _point_record(old_parts[i], gtype, "removed", "part",
                            part_old=i + 1)
        if rec:
            points.append(rec)

    n_modified = 0
    for i, j in modified:
        d = geometry_diff(old_parts[i], new_parts[j], tol, gtype)
        if not d["changed"]:
            continue
        n_modified += 1
        pieces = []
        if d["added"]:
            added_geoms.append(d["added"])
            pieces.append(d["added"])
        if d["removed"]:
            removed_geoms.append(d["removed"])
            pieces.append(d["removed"])
        where = QgsGeometry.collectGeometry(pieces) if pieces else new_parts[j]
        rec = _point_record(where, gtype, "modified", "part",
                            part_old=i + 1, part_new=j + 1)
        if rec:
            rec["measure"] = d["added_m"] + d["removed_m"]
            points.append(rec)

    def merge(geoms):
        if not geoms:
            return None
        merged = QgsGeometry.collectGeometry(geoms)
        return merged if _geom_present(merged) else None

    added, removed = merge(added_geoms), merge(removed_geoms)
    old_m = _geom_measure(old_g, gtype)
    added_m = sum(_geom_measure(g, gtype) for g in added_geoms)
    removed_m = sum(_geom_measure(g, gtype) for g in removed_geoms)
    return {
        "changed": bool(added or removed),
        "added": added,
        "removed": removed,
        "added_m": added_m,
        "removed_m": removed_m,
        "pct": ((added_m + removed_m) / old_m * 100.0) if old_m else None,
        "parts_old": len(old_parts),
        "parts_new": len(new_parts),
        "points": points,
        "parts_changed": len(only_new) + len(only_old) + n_modified,
    }


def analyse_geometry(old_g, new_g, tol, gtype, by_part):
    """Compare one feature's geometry in OLD and NEW.

    *by_part* False: the whole (multi)geometry is compared as one shape.
    *by_part* True : every part is matched and compared individually.
    """
    old_ok, new_ok = _geom_present(old_g), _geom_present(new_g)
    if old_ok and new_ok:
        if by_part:
            return geometry_diff_parts(old_g, new_g, tol, gtype)
        gd = geometry_diff(old_g, new_g, tol, gtype)
        gd["points"] = _region_points(gd, gtype) if gd["changed"] else []
        return gd
    gd = geometry_diff(old_g, new_g, tol, gtype)
    gd["points"] = _feature_points(gd, gtype) if (old_ok or new_ok) else []
    return gd


def _fmt_num(v):
    text = f"{v:.6f}".rstrip("0").rstrip(".")
    return text if text not in ("", "-0") else "0"


class _SortItem(QTableWidgetItem):
    """Table item that sorts numbers numerically (10 after 2)."""

    @staticmethod
    def _key(text):
        try:
            return (0, float(text.replace(",", "").rstrip("%")), "")
        except ValueError:
            return (1, 0.0, text.lower())

    def __lt__(self, other):
        return self._key(self.text()) < self._key(other.text())


# ============================================================================
# THEME / ICONS (shared Attribute Table Studio look - the Compare palette)
# ============================================================================

def _qgis_icon(theme_name, *_fallback):
    """Built-in QGIS theme icon."""
    return icons.native(theme_name)


def _field_icon(field, theme):
    return icons.field_for(field, theme)


def _field_tooltip(field):
    return W.field_tooltip(field)


class _CheckLabel(QLabel):
    """Text box that toggles its checkbox when clicked (acts as the legend)."""

    def __init__(self, text, checkbox, parent=None):
        super().__init__(text, parent)
        self._checkbox = checkbox
        self.setCursor(Qt.CursorShape.PointingHandCursor)
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)

    def mouseReleaseEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self._checkbox.toggle()
        super().mouseReleaseEvent(event)


class _CompareDelegate(QStyledItemDelegate):
    """Draws a changed value that is still pending as  OLD | NEW : two equal
    halves with a divider, so the dividers line up down the column.  Every
    other cell (and an accepted / rejected value) is drawn normally."""

    PAD = 6

    def __init__(self, theme, parent=None):
        super().__init__(parent)
        self._divider = QColor(theme.accent_of("ChangedField"))

    @staticmethod
    def pair(index):
        if index.data(_ROLE_KIND) != "changed" or index.data(_ROLE_DECISION):
            return None
        pair = index.data(_ROLE_USER)
        if isinstance(pair, (tuple, list)) and len(pair) == 2:
            return pair
        return None

    def paint(self, painter, option, index):
        pair = self.pair(index)
        if pair is None:
            super().paint(painter, option, index)
            return
        opt = QStyleOptionViewItem(option)
        self.initStyleOption(opt, index)
        rect = opt.rect
        painter.save()
        if opt.state & QStyle.StateFlag.State_Selected:
            painter.fillRect(rect, opt.palette.highlight())
            fg = opt.palette.highlightedText().color()
            line = fg
        else:
            bg = index.data(Qt.ItemDataRole.BackgroundRole)
            if bg is not None:
                painter.fillRect(rect, bg)
            brush = index.data(Qt.ItemDataRole.ForegroundRole)
            fg = brush.color() if brush is not None else opt.palette.text().color()
            line = self._divider
        mid = rect.left() + rect.width() // 2
        painter.setPen(QPen(line, 1))
        painter.drawLine(mid, rect.top() + 3, mid, rect.bottom() - 3)
        painter.setPen(fg)
        painter.setFont(opt.font)
        fm = opt.fontMetrics
        pad = self.PAD
        halves = (QRect(rect.left() + pad, rect.top(), mid - rect.left() - 2 * pad,
                        rect.height()),
                  QRect(mid + pad, rect.top(), rect.right() - mid - 2 * pad + 1,
                        rect.height()))
        for area, text in zip(halves, pair):
            painter.drawText(
                area, Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                fm.elidedText(str(text), Qt.TextElideMode.ElideRight, area.width()))
        painter.restore()

    def sizeHint(self, option, index):
        hint = super().sizeHint(option, index)
        pair = self.pair(index)
        if pair is not None:
            fm = option.fontMetrics
            half = max(fm.horizontalAdvance(str(v)) for v in pair)
            hint.setWidth(2 * (half + 2 * self.PAD) + 1)
        return hint


# ============================================================================
# MAIN PANEL
# ============================================================================

class GeoTableCompareDialog(QWidget):
    """
    Main comparison panel for the Geometry and Table Compare tool.

    Directional comparison:  OLD (Baseline) ──→ NEW (Updated)
    - Added     = feature exists only in NEW
    - Deleted   = feature exists only in OLD
    - Modified  = feature changed between OLD → NEW
    - Unchanged = feature is identical in both

    Shown by the Attribute Table Studio window / dock host (see ui/tool_host.py).
    """

    ACTION_SIZE = QSize(190, 28)    # Validate Data and Compare, as in v1.1

    def __init__(self, iface, parent=None):
        super().__init__(parent)
        self.setObjectName("gtsPanel")
        self.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, True)
        self.iface = iface
        self.theme = Theme()
        self.metrics = Metrics(self.font())
        # badges / "Join field" / "Compare geometry" share one column width
        fm, bold = self.metrics.fm, self.metrics.bold_fm
        self.LABEL_COLUMN = max(
            150, bold.horizontalAdvance("NEW  Updated") + 30,
            fm.horizontalAdvance("Compare geometry") + 34)
        # the width comes from the layout: the Results and Review rows never
        # get narrower than their buttons and legends need
        self.setMinimumHeight(380)

        # Data storage
        self.decisions = {}            # join id -> row decision
        self.cell_decisions = {}       # (join id, field) -> value decision
        self._fid_maps = {"old": {}, "new": {}}   # join id -> feature id
        self._run_layer_ids = (None, None)
        self._help_dialog = None
        self.columns_to_check = []
        self._compare_info = {}
        self._old_fields = None
        self._geom_diffs = {}
        self._geom_crs = None
        self._geom_gtype = None
        self._map_layer_ids = []
        self._part_cache = {}
        self._mode_choice = "parts"       # user's last pick; default: split
        self._review_saved = ""
        self._review_layer_id = None
        self._review_note = ""
        self._geom_by_part_run = False
        self._reviewable_total = 0
        self._value_total = 0            # changed values (OLD | NEW cells)
        self._has_results = False

        self._build_ui()
        self._apply_style()
        self._wire_signals()
        self._on_layer_changed()

    # host contract (ui/tool_host.py) -----------------------------------
    def window_title(self):
        return PLUGIN_NAME

    def full_window_title(self):
        """Compare keeps its own title (no suite prefix)."""
        return PLUGIN_NAME

    def shutdown(self):
        """Nothing runs in the background; layers added to the map stay."""

    # ====================================================================
    # UI CONSTRUCTION
    # ====================================================================

    def _build_ui(self):
        m = self.metrics
        root = QVBoxLayout(self)
        root.setContentsMargins(m.margin, m.margin, m.margin, m.margin)
        root.setSpacing(8)

        root.addWidget(self._create_input_card())
        root.addSpacing(2)
        root.addWidget(self._create_summary_bar())
        root.addWidget(self._create_review_bar())

        self.results_stack = QStackedWidget()
        self.results_stack.addWidget(self._create_empty_state())
        self.results_stack.addWidget(self._create_results_table())
        root.addWidget(self.results_stack, 1)

        root.addWidget(self._create_bottom_bar())

    # ------------------------------------------------------------------
    def _create_input_card(self):
        m, t = self.metrics, self.theme
        card = QFrame()
        card.setObjectName("gtsCard")
        grid = QGridLayout(card)
        grid.setContentsMargins(14, 12, 14, 12)
        grid.setHorizontalSpacing(10)
        grid.setVerticalSpacing(8)
        grid.setColumnStretch(1, 1)
        grid.setColumnMinimumWidth(0, self.LABEL_COLUMN)

        # ---- OLD / NEW layer rows ----
        self.old_table_combo = self._new_layer_combo()
        self.new_table_combo = self._new_layer_combo()
        self.old_info = QLabel("")
        self.new_info = QLabel("")
        for lbl in (self.old_info, self.new_info):
            lbl.setObjectName("gtsDim")
            lbl.setMinimumWidth(220)

        grid.addWidget(self._badge("OLD", "Baseline", t.old_colour), 0, 0)
        grid.addWidget(self.old_table_combo, 0, 1)
        grid.addWidget(self.old_info, 0, 2)
        grid.addWidget(self._badge("NEW", "Updated", t.new_colour), 1, 0)
        grid.addWidget(self.new_table_combo, 1, 1)
        grid.addWidget(self.new_info, 1, 2)

        self.swap_btn = W.primary_square(m, icons.native(icons.SWAP),
                                         "Swap OLD and NEW layers")
        grid.addWidget(self.swap_btn, 0, 3, 2, 1,
                       Qt.AlignmentFlag.AlignVCenter | Qt.AlignmentFlag.AlignRight)
        # same place as in every Attribute Table Studio tool: top right
        self.dock_btn = W.dock_button(m)
        grid.addWidget(self.dock_btn, 0, 4,
                       Qt.AlignmentFlag.AlignTop | Qt.AlignmentFlag.AlignRight)

        # ---- Join field + column selector ----
        join_lbl = QLabel("Join field")
        join_lbl.setFixedWidth(self.LABEL_COLUMN)
        grid.addWidget(join_lbl, 2, 0)

        self.join_field_combo = QComboBox()
        self.join_field_combo.setMinimumWidth(240)
        self.join_field_combo.setMaximumWidth(340)
        self.join_field_combo.setFixedHeight(m.control_h)
        self.join_field_combo.setIconSize(QSize(m.icon, m.icon))
        self.join_field_combo.setMaxVisibleItems(16)
        self.join_field_combo.view().setMinimumWidth(320)
        self.join_field_combo.setToolTip(
            "Unique identifier field common to both layers.\n"
            "Used to match rows between OLD and NEW."
        )

        self.column_filter_btn = W.button(
            m, "Field(s) to Compare", icons.native(icons.COLUMNS),
            "Choose which fields are compared\nwhen detecting modified features")

        # ---- Geometry comparison ----
        self.geom_compare_cb = QCheckBox("Compare geometry")
        self.geom_compare_cb.setFixedWidth(self.LABEL_COLUMN)
        self.geom_compare_cb.setToolTip(
            "Compare geometry properties between OLD and NEW.\n"
            "Adds virtual fields to the comparison results."
        )
        grid.addWidget(self.geom_compare_cb, 3, 0)

        self.geom_length_cb = QCheckBox("Shape Length")
        self.geom_area_cb = QCheckBox("Shape Area")
        self.geom_centroid_cb = QCheckBox("Centroid XY")
        self.geom_vertices_cb = QCheckBox("Vertex Count")
        geom_row = W.FlowLayout(h_spacing=14, v_spacing=6)
        for cb in self._geom_checkboxes():
            cb.setEnabled(False)
            cb.setFixedHeight(m.control_h)
            geom_row.addWidget(cb)

        self.geom_tol_spin = QDoubleSpinBox()
        self.geom_tol_spin.setDecimals(6)
        self.geom_tol_spin.setRange(0.0, 1000000.0)
        self.geom_tol_spin.setSingleStep(0.001)
        self.geom_tol_spin.setValue(DEFAULT_GEOM_TOLERANCE)
        self.geom_tol_spin.setFixedWidth(100)
        self.geom_tol_spin.setFixedHeight(m.control_h)
        tol_tip = (
            "Geometry that moved by less than this distance (in the layer's\n"
            "map units) is treated as unchanged. Use 0 for an exact match."
        )
        tol_box = W.labelled("Tolerance", self.geom_tol_spin)
        self.geom_tol_label = tol_box.label
        self.geom_tol_label.setToolTip(tol_tip)
        self.geom_tol_spin.setToolTip(tol_tip)

        # How a multipart geometry is compared (auto-detected from the input)
        self.geom_mode_combo = QComboBox()
        self.geom_mode_combo.setFixedHeight(m.control_h)
        self.geom_mode_combo.setFixedWidth(130)
        self.geom_mode_combo.setIconSize(QSize(m.icon, m.icon))
        self._add_mode_item("Singlepart", "whole")
        mode_box = W.labelled("Method", self.geom_mode_combo)
        self.geom_mode_label = mode_box.label
        self.geom_mode_label.setToolTip(MODE_TIP)
        self.geom_mode_combo.setToolTip(MODE_TIP)

        # Where the review points are stored (temporary layer or a file)
        self.review_out_combo = QComboBox()
        self.review_out_combo.setFixedHeight(m.control_h)
        self.review_out_combo.setFixedWidth(210)
        self.review_out_combo.setIconSize(QSize(m.icon, m.icon))
        for text, key, icon in (
                ("Temporary Points", "temp", "/mIconPointLayer.svg"),
                ("Save as shapefile (.shp)", "shp",
                 "/mActionNewVectorLayer.svg"),
                ("Save as geopackage (.gpkg)", "gpkg",
                 "/mActionNewGeoPackageLayer.svg")):
            self.review_out_combo.addItem(_qgis_icon(icon), text, key)
        review_box = W.labelled("Review points", self.review_out_combo)
        self.review_out_label = review_box.label
        self._review_path = ""
        self._review_fmt = "temp"
        self._set_review_tooltip()

        for box in (tol_box, mode_box, review_box):
            geom_row.addWidget(box)

        for w in (self.geom_tol_label, self.geom_tol_spin,
                  self.geom_mode_label, self.geom_mode_combo,
                  self.review_out_label, self.review_out_combo):
            w.setEnabled(False)
        grid.addLayout(geom_row, 3, 1, 1, 4)

        # ---- Help, clear validation and the action buttons ----
        self.help_btn = W.primary_square(
            m, icons.native(icons.HELP),
            f"What {PLUGIN_NAME} is for and how to use it")
        self.clear_validation_btn = W.primary_square(
            m, icons.native(icons.CLEAR), "Clear the validation results")
        self.clear_validation_btn.setEnabled(False)

        # identical size, the v1.1 buttons
        self.validate_btn = QPushButton(" Validate Data")
        self.validate_btn.setIcon(icons.action_icon("validate", t.text))
        self.validate_btn.setToolTip(
            "Check for duplicate IDs, field mismatches,\n"
            "and null values before comparing")
        self.compare_button = QPushButton("Compare OLD → NEW")
        self.compare_button.setProperty("role", "primary")
        self.compare_button.setToolTip(
            "Detect added, deleted and modified features")
        # v1.1 size; wider only if the bold 10 pt label would not fit
        bold = QFont(self.compare_button.font())
        bold.setPointSizeF(10)
        bold.setBold(True)
        label_w = QFontMetrics(bold).horizontalAdvance(" Validate Data") + 18
        label_w = max(label_w, QFontMetrics(bold).horizontalAdvance(
            self.compare_button.text()))
        size = QSize(max(self.ACTION_SIZE.width(), label_w + 2 * 14 + 10),
                     self.ACTION_SIZE.height())
        for b in (self.validate_btn, self.compare_button):
            b.setObjectName("actionBtn")
            b.setFixedSize(size)
            b.setIconSize(QSize(18, 18))
            b.setCursor(Qt.CursorShape.PointingHandCursor)

        # Join row: field pickers on the left, action buttons on the right
        join_row = QHBoxLayout()
        join_row.setSpacing(m.gap)
        join_row.addWidget(self.join_field_combo)
        fields_row = W.ButtonRow(m)              # icon-only when narrow
        self.column_filter_btn.setProperty("gtsNameTip", True)
        fields_row.add(self.column_filter_btn)
        fields_row.add_stretch()
        join_row.addWidget(fields_row, 1)
        join_row.addWidget(self.help_btn)
        join_row.addWidget(self.clear_validation_btn)
        join_row.addSpacing(m.gap)
        join_row.addWidget(self.validate_btn)
        join_row.addWidget(self.compare_button)
        grid.addLayout(join_row, 2, 1, 1, 4)

        # ---- Validation results (hidden by default) ----
        self.validation_frame = QFrame()
        self.validation_frame.setObjectName("validation")
        vf = QVBoxLayout(self.validation_frame)
        vf.setContentsMargins(10, 8, 10, 8)
        self.validation_label = QLabel("")
        self.validation_label.setWordWrap(True)
        self.validation_label.setTextFormat(Qt.TextFormat.RichText)
        self.validation_label.setTextInteractionFlags(
            Qt.TextInteractionFlag.TextSelectableByMouse)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.Shape.NoFrame)
        scroll.setMaximumHeight(130)
        scroll.setWidget(self.validation_label)
        scroll.setObjectName("validationScroll")
        vf.addWidget(scroll)
        self.validation_scroll = scroll
        self.validation_frame.setVisible(False)
        grid.addWidget(self.validation_frame, 4, 0, 1, 5)

        return card

    def _geom_checkboxes(self):
        return (self.geom_length_cb, self.geom_area_cb,
                self.geom_centroid_cb, self.geom_vertices_cb)

    def _new_layer_combo(self):
        combo = W.LayerCombo(self.metrics)
        combo.setFilters(_VECTOR_FILTER)
        combo.setAllowEmptyLayer(True)
        combo.setMinimumWidth(240)
        return combo

    def _badge(self, title, subtitle, colour):
        lbl = QLabel(f"<span style='color:{colour}; font-weight:700;'>{title}"
                     f"</span>&nbsp;&nbsp;{subtitle}")
        lbl.setObjectName("gtsBadge")
        lbl.setTextFormat(Qt.TextFormat.RichText)
        lbl.setFixedSize(self.LABEL_COLUMN, self.metrics.control_h)
        return lbl

    def _row_title(self, text):
        lbl = W.label(text, "title")
        fm = self.metrics.bold_fm
        lbl.setFixedWidth(max(fm.horizontalAdvance("Results"),
                              fm.horizontalAdvance("Review")) + 22)
        return lbl

    # ------------------------------------------------------------------
    def _create_summary_bar(self):
        m = self.metrics
        bar = QWidget()
        lay = QHBoxLayout(bar)
        lay.setContentsMargins(2, 0, 0, 0)
        lay.setSpacing(m.gap)
        lay.addWidget(self._row_title("Results"))

        # v1.1 legend: tick box + rounded box in the row colour + live count
        self.filter_buttons = {}
        self.filter_boxes = {}
        self.filter_counts = {}
        labels = {status: status for status in STATUSES}
        tips = {
            "Added": "Features that exist only in the NEW version",
            "Deleted": "Features that exist in OLD but are missing from NEW",
            "Modified": "Features whose compared values changed "
                        "between OLD and NEW",
            "Unchanged": "Features identical in both OLD and NEW",
        }
        for status in STATUSES:
            cb = QCheckBox()
            cb.setObjectName(f"filter{status}")
            cb.setChecked(True)
            cb.setToolTip(f"{tips[status]}\n(tick to show / untick to hide)")
            cb.setFixedHeight(m.control_h)
            box = _CheckLabel(labels[status], cb)
            box.setObjectName(f"filterBox{status}")
            box.setToolTip(cb.toolTip())
            box.setFixedHeight(m.control_h)
            box.setMinimumWidth(
                m.bold_fm.horizontalAdvance(labels[status]) + 34)
            cb.toggled.connect(
                lambda _checked, s=status: self._style_filter_box(s))
            count = QLabel("")
            count.setObjectName("gtsDim")
            self.filter_buttons[status] = cb
            self.filter_boxes[status] = box
            self.filter_counts[status] = count
            item = QHBoxLayout()
            item.setSpacing(m.gap)
            item.addWidget(cb)
            item.addWidget(box)
            item.addWidget(count)
            lay.addLayout(item)
            lay.addSpacing(m.gap)
            self._style_filter_box(status)
        lay.addStretch()

        # ---- Filter by field value (field picker + suggestions) ----
        self.filter_field_combo = QComboBox()
        self.filter_field_combo.setFixedHeight(m.control_h)
        field_w = m.fm.horizontalAdvance("All fields") + m.icon + 44
        self.filter_field_combo.setFixedWidth(field_w)
        self.filter_field_combo.setIconSize(QSize(m.icon, m.icon))
        self.filter_field_combo.setMaxVisibleItems(16)
        self.filter_field_combo.view().setMinimumWidth(280)
        self.filter_field_combo.addItem(self._all_fields_icon(), "All fields", -1)
        self.filter_field_combo.setToolTip(
            "Choose the field the filter text is matched against")
        lay.addWidget(self.filter_field_combo)

        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("Filter by value")
        self.search_edit.setClearButtonEnabled(True)
        self.search_edit.setFixedHeight(m.control_h)
        # shrinks with the panel, never below the field picker beside it
        self.search_edit.setMinimumWidth(field_w)
        self.search_edit.setMaximumWidth(300)
        self.search_edit.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.search_edit.setToolTip(
            "Type a value to show only matching rows.\n"
            "Suggestions come from the selected field.")
        self.value_model = QStringListModel(self)
        self.value_completer = QCompleter(self.value_model, self)
        self.value_completer.setCaseSensitivity(
            Qt.CaseSensitivity.CaseInsensitive)
        self.value_completer.setFilterMode(Qt.MatchFlag.MatchContains)
        self.value_completer.setMaxVisibleItems(12)
        self.search_edit.setCompleter(self.value_completer)
        lay.addWidget(self.search_edit, 1)

        export_menu = QMenu(self)
        self.export_xlsx_action = export_menu.addAction(
            _qgis_icon(icons.XLSX), "Excel (.xlsx) — coloured, as shown")
        self.export_csv_action = export_menu.addAction(
            _qgis_icon(icons.CSV), "CSV (.csv) — plain values")
        if not HAS_OPENPYXL:
            self.export_xlsx_action.setEnabled(False)
            self.export_xlsx_action.setToolTip(
                "Excel export needs the 'openpyxl' Python package")
        self.export_btn = W.menu_square(
            m, icons.native(icons.EXPORT),
            "Export\nThe rows shown in the table to Excel or CSV", export_menu)
        self.export_btn.setEnabled(False)
        lay.addWidget(self.export_btn)
        return bar

    def _set_filter_count(self, status, count):
        self.filter_counts[status].setText(f"· {count:,}")

    def _style_filter_box(self, status):
        """Coloured legend box next to a filter tick (dimmed when unticked)."""
        t = self.theme
        box = self.filter_boxes[status]
        if self.filter_buttons[status].isChecked():
            box.setStyleSheet(
                f"background: {t.bg(status).name()}; color: {t.row_text}; "
                f"border: 1px solid {t.accent_of(status).name()}; "
                f"border-radius: 4px; padding: 0 10px; font-weight: bold;")
        else:
            box.setStyleSheet(
                f"background: transparent; color: {t.text_dim}; "
                f"border: 1px solid {t.border}; border-radius: 4px; "
                f"padding: 0 10px;")

    def _all_fields_icon(self):
        return _qgis_icon(icons.TABLE)

    def _field_icon_and_tip(self, name):
        """Icon + tooltip for a result field (layer field or geometry value)."""
        fields = self._old_fields
        if fields is not None and fields.indexOf(name) >= 0:
            fld = fields.field(name)
            return _field_icon(fld, self.theme), _field_tooltip(fld)
        if name == "_geom_status":
            return (icons.field(FK.TEXT, self.theme),
                    f"<b>{name}</b><br>Exact geometry comparison result")
        integer = name in ('_vertex_count', '_parts', '_parts_changed')
        icon = icons.field(FK.INTEGER if integer else FK.REAL, self.theme)
        return icon, (f"<b>{name}</b><br>Geometry property "
                      f"({'integer' if integer else 'decimal'}, calculated)")

    def _create_review_bar(self):
        m = self.metrics
        bar = QWidget()
        lay = QHBoxLayout(bar)
        lay.setContentsMargins(2, 0, 0, 0)
        lay.setSpacing(m.gap)

        title = self._row_title("Review")
        title.setToolTip(
            "Mark Added / Modified rows as Accepted or Rejected.\n"
            "Nothing is written to your layers — decisions are saved "
            "in the exported CSV\n(a Rejected row exports its OLD value)."
        )
        lay.addWidget(title)

        self.review_scope = QComboBox()
        self.review_scope.addItems(["Selected rows", "All visible rows"])
        self.review_scope.setFixedHeight(m.control_h)
        self.review_scope.setMinimumWidth(150)
        self.review_scope.setToolTip("Which rows the buttons apply to")
        lay.addWidget(self.review_scope)

        # the buttons turn icon-only when the panel is too narrow for them
        buttons = W.ButtonRow(m)
        self.accept_btn = self._review_button(
            "Accept", icons.SUCCESS, "accept",
            "Mark the Added / Modified rows as accepted")
        self.reject_btn = self._review_button(
            "Reject", icons.DELETE, "reject",
            "Mark the Added / Modified rows as rejected (exports the OLD value)")
        self.reset_btn = self._review_button(
            "Reset", icons.UNDO, "reset", "Return the rows to Pending")
        for b in (self.accept_btn, self.reject_btn, self.reset_btn):
            buttons.add(b)

        self.zoom_btn = W.button(
            m, "Zoom to selection", icons.native(icons.ZOOM_SELECTED),
            "Select the features of the selected rows in the OLD and NEW\n"
            "layers, zoom the map to them and flash them\n"
            "(the NEW shape; the OLD one for Deleted rows).")
        self.zoom_btn.setEnabled(False)
        self.map_btn = W.button(
            m, "Show on map", icons.native(icons.GEOM_DIFF),
            "Add the geometry that was ADDED (green) and REMOVED (red) for the\n"
            "selected rows (or all visible rows if none are selected) to the\n"
            "map, and zoom to it.\n"
            "Needs 'Compare geometry' to be ticked when comparing.")
        self.map_btn.setEnabled(False)
        buttons.add_spacing(m.gap)
        for b in (self.zoom_btn, self.map_btn):
            b.setProperty("gtsNameTip", True)
            buttons.add(b)
        buttons.add_stretch()
        self.review_buttons = buttons
        lay.addWidget(buttons, 1)
        self.changed_field_legend = QLabel("Changed")
        self.changed_field_legend.setObjectName("legendChanged")
        self.changed_field_legend.setFixedHeight(m.compact_h)
        self.changed_field_legend.setToolTip(
            "A highlighted cell shows a value that differs between OLD and "
            "NEW:\nOLD value on the left | NEW value on the right "
            "(blank = empty / NULL).\nRight-click it to accept or reject it."
        )
        lay.addWidget(self.changed_field_legend)
        self.accepted_legend = QLabel("Accepted")
        self.accepted_legend.setObjectName("legendAccepted")
        self.accepted_legend.setToolTip(
            "An accepted change: the NEW value is kept (and exported)")
        self.rejected_legend = QLabel("Rejected")
        self.rejected_legend.setObjectName("legendRejected")
        self.rejected_legend.setToolTip(
            "A rejected change: the OLD value is kept (and exported)")
        for legend in (self.accepted_legend, self.rejected_legend):
            legend.setFixedHeight(m.compact_h)
            lay.addWidget(legend)
        self.geom_summary_label = W.ElideLabel()
        self.geom_summary_label.setObjectName("gtsDim")
        self.geom_summary_label.setMinimumWidth(60)
        self.geom_summary_label.setSizePolicy(QSizePolicy.Policy.Preferred,
                                              QSizePolicy.Policy.Preferred)
        lay.addWidget(self.geom_summary_label)

        self.review_widgets = (
            self.review_scope, self.accept_btn, self.reject_btn,
            self.reset_btn, self.zoom_btn)
        for w in self.review_widgets:
            w.setEnabled(False)
        return bar

    def _review_button(self, text, icon_name, role, tip):
        btn = W.button(self.metrics, text, icons.native(icon_name), tip)
        btn.setProperty("role", role)
        btn.setProperty("gtsNameTip", True)       # icon-only: name in the tooltip
        btn.setMinimumWidth(96)
        return btn

    def _create_empty_state(self):
        page = QFrame()
        page.setObjectName("emptyPage")
        lay = QVBoxLayout(page)
        self.empty_label = QLabel(
            f"<div style='font-size:{self.metrics.display_pt}pt; font-weight:600;'>"
            "No comparison yet</div><br>"
            "1.&nbsp; Pick the <b>OLD</b> and <b>NEW</b> layers<br>"
            "2.&nbsp; Choose the <b>Join field</b> (unique ID)<br>"
            "3.&nbsp; Press <b>Compare OLD → NEW</b>"
        )
        self.empty_label.setObjectName("emptyState")
        self.empty_label.setTextFormat(Qt.TextFormat.RichText)
        self.empty_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lay.addWidget(self.empty_label)
        return page

    def _create_results_table(self):
        self.results_table = QTableWidget()
        tbl = self.results_table
        tbl.setObjectName("gtsTable")
        tbl.setHorizontalHeader(W.SortHeader(self.theme, tbl))
        tbl.setSortingEnabled(True)
        # single cells can be picked, so values are accepted one by one;
        # the row header still selects whole rows
        tbl.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectItems)
        tbl.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        tbl.setItemDelegate(_CompareDelegate(self.theme, tbl))
        tbl.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        tbl.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        tbl.setAlternatingRowColors(False)
        tbl.setShowGrid(True)
        tbl.setWordWrap(False)
        tbl.horizontalHeader().setDefaultAlignment(
            Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)
        tbl.horizontalHeader().setMinimumHeight(self.metrics.font_h + 14)
        tbl.horizontalHeader().sectionClicked.connect(self._on_column_sort)
        return tbl

    def _create_bottom_bar(self):
        """Full-width progress bar (shown while comparing) above the shared
        status strip - the v1.1 bottom bar."""
        widget = QWidget()
        layout = QVBoxLayout(widget)
        layout.setContentsMargins(0, 4, 0, 0)
        layout.setSpacing(4)
        self.progress_bar = QProgressBar()
        self.progress_bar.setObjectName("gtsProgressWide")
        self.progress_bar.setTextVisible(True)
        self.progress_bar.setFormat("%p%")
        self.progress_bar.setFixedHeight(self.metrics.font_h + 4)
        self.progress_bar.setValue(0)
        self.progress_bar.setVisible(False)
        layout.addWidget(self.progress_bar)
        self.status = W.StatusStrip(self.metrics)
        self.rows_label = self.status.extra
        # row decisions sit beside the row count, on the right
        self.review_tally = QLabel("")
        self.review_tally.setObjectName("gtsDim")
        strip = self.status.layout()
        strip.insertWidget(strip.indexOf(self.rows_label), self.review_tally)
        strip.insertSpacing(strip.indexOf(self.rows_label), self.metrics.group_gap)
        layout.addWidget(self.status)
        return widget

    # ------------------------------------------------------------------
    # Styling
    # ------------------------------------------------------------------
    def _apply_style(self):
        t, m = self.theme, self.metrics
        self.setStyleSheet(stylesheet(t, m) + f"""
            QFrame#emptyPage {{ background: {t.base};
                border: 1px solid {t.border}; border-radius: 4px; }}
            QFrame#validation {{ background: {t.base};
                border: 1px solid {t.border}; border-radius: 4px; }}
            QScrollArea#validationScroll,
            QScrollArea#validationScroll > QWidget > QWidget
                {{ background: transparent; }}
            QLabel#emptyState {{ color: {t.text_dim}; }}
            QPushButton#actionBtn {{ background: {t.card}; color: {t.text};
                border: 1px solid {t.border}; border-radius: 4px;
                font-size: 10pt; font-weight: 700; padding: 0 14px; }}
            QPushButton#actionBtn:hover {{ background: {t.card_hover}; }}
            QPushButton#actionBtn:disabled {{ color: {t.text_dim}; }}
            QPushButton#actionBtn[role="primary"] {{ background: {t.accent};
                color: white; border: none; }}
            QPushButton#actionBtn[role="primary"]:hover {{
                background: {t.accent_hover}; }}
            QPushButton#actionBtn[role="primary"]:disabled {{
                background: {t.disabled_bg}; color: {t.disabled_fg}; }}
            QLabel#gtsBadge {{ background: {t.base}; color: {t.text_strong};
                border: 1px solid {t.border}; border-radius: 4px;
                padding: 0px 10px; }}
            QLabel#legendChanged {{ background: {t.bg('ChangedField').name()};
                color: {t.row_text};
                border: 1px solid {t.accent_of('ChangedField').name()};
                border-radius: 4px; padding: 0px 8px; }}
            QLabel#legendAccepted {{ background: {t.bg('AcceptedValue').name()};
                color: {t.row_text};
                border: 1px solid {t.accent_of('AcceptedValue').name()};
                border-radius: 4px; padding: 0px 8px; }}
            QLabel#legendRejected {{ background: {t.bg('RejectedValue').name()};
                color: {t.row_text};
                border: 1px solid {t.accent_of('RejectedValue').name()};
                border-radius: 4px; padding: 0px 8px; }}
            QToolButton#gtsButton[role="accept"] {{
                color: {t.accent_of('AcceptedValue').name()};
                border-color: {t.accent_of('AcceptedValue').name()}; }}
            QToolButton#gtsButton[role="reject"] {{
                color: {t.accent_of('RejectedValue').name()};
                border-color: {t.accent_of('RejectedValue').name()}; }}
            QToolButton#gtsButton[role="accept"]:hover {{
                background: {t.accepted_bg.name()}; color: {t.row_text}; }}
            QToolButton#gtsButton[role="reject"]:hover {{
                background: {t.rejected_bg.name()}; color: {t.row_text}; }}
            QToolButton#gtsButton[role="accept"]:disabled,
            QToolButton#gtsButton[role="reject"]:disabled {{
                color: {t.text_faint}; border-color: {t.border_soft}; }}
        """)

        vh = self.results_table.verticalHeader()
        vh.setDefaultSectionSize(m.font_h + 10)
        vh.setMinimumWidth(40)

        # Selection colour via the palette: a ::item:selected style rule makes
        # Qt re-lay-out the selected row (text shifts, inner focus box appears)
        pal = self.results_table.palette()
        pal.setColor(QPalette.ColorRole.Highlight, QColor(t.selected))
        pal.setColor(QPalette.ColorRole.HighlightedText,
                     QColor(t.selected_text))
        self.results_table.setPalette(pal)

    def _set_status(self, text, level="info"):
        """Update the status strip message and its coloured accent."""
        self.status.set_status(text, level)

    # ====================================================================
    # SIGNAL WIRING
    # ====================================================================

    def _wire_signals(self):
        self.old_table_combo.layerChanged.connect(self._on_layer_changed)
        self.new_table_combo.layerChanged.connect(self._on_layer_changed)
        self.join_field_combo.currentIndexChanged.connect(
            self._update_ready_state)
        self.swap_btn.clicked.connect(self._swap_layers)

        self.geom_compare_cb.toggled.connect(self._on_geom_toggle)

        self.validate_btn.clicked.connect(self.validate_data)
        self.help_btn.clicked.connect(self.show_help)
        self.clear_validation_btn.clicked.connect(self.clear_validation)
        self.compare_button.clicked.connect(self.compare_tables)
        self.column_filter_btn.clicked.connect(self.select_columns_to_check)
        self.export_xlsx_action.triggered.connect(self.export_xlsx)
        self.export_csv_action.triggered.connect(self.export_results)

        for btn in self.filter_buttons.values():
            btn.toggled.connect(self.apply_filters)
        self.filter_field_combo.currentIndexChanged.connect(
            self._on_filter_field_changed)
        self.search_edit.textChanged.connect(self.apply_filters)

        self.accept_btn.clicked.connect(lambda: self._apply_decision("Accepted"))
        self.reject_btn.clicked.connect(lambda: self._apply_decision("Rejected"))
        self.reset_btn.clicked.connect(lambda: self._apply_decision(None))
        self.map_btn.clicked.connect(self.show_changes_on_map)
        self.zoom_btn.clicked.connect(self.zoom_to_selection)
        self.results_table.customContextMenuRequested.connect(self._on_table_menu)
        self.review_out_combo.activated.connect(self._on_review_out_activated)
        self.geom_mode_combo.activated.connect(self._on_mode_activated)
        self.results_table.cellDoubleClicked.connect(self._on_row_double_click)

    # ====================================================================
    # LAYER & FIELD MANAGEMENT
    # ====================================================================

    def _swap_layers(self):
        old_layer = self.old_table_combo.currentLayer()
        new_layer = self.new_table_combo.currentLayer()
        self.old_table_combo.setLayer(new_layer)
        self.new_table_combo.setLayer(old_layer)

    def _on_layer_changed(self, *_args):
        self._update_join_fields()
        self._update_geom_options()
        self._update_geom_mode()
        self._update_layer_info()
        self.validation_frame.setVisible(False)
        self.clear_validation_btn.setEnabled(False)
        self._update_ready_state()

    # ------------------------------------------------------------------
    # Multipart detection + analysis mode
    # ------------------------------------------------------------------
    def _layer_max_parts(self, layer):
        """Most parts found in one feature (1 for single-part / non-spatial)."""
        if layer is None or not layer.isSpatial():
            return 1
        key = (layer.id(), layer.featureCount())
        if key not in self._part_cache:
            most = 1
            if QgsWkbTypes.isMultiType(layer.wkbType()):
                request = QgsFeatureRequest()
                request.setNoAttributes()
                request.setLimit(MAX_PART_SCAN)
                for feat in layer.getFeatures(request):
                    most = max(most, _geom_parts(feat.geometry()))
            self._part_cache[key] = most
        return self._part_cache[key]

    def _add_mode_item(self, text, key):
        icon_name = ("/algorithms/mAlgorithmMultiToSingle.svg"
                     if text == "Singlepart"
                     else "/algorithms/mAlgorithmSingleToMulti.svg")
        self.geom_mode_combo.addItem(
            _qgis_icon(icon_name, "dot", self.theme.text), text, key)

    def _update_geom_mode(self):
        """Auto-detect multipart input and offer the matching choices."""
        on = self.geom_compare_cb.isChecked()
        for w in (self.geom_tol_label, self.geom_tol_spin,
                  self.geom_mode_label, self.review_out_label,
                  self.review_out_combo):
            w.setEnabled(on)

        most_parts = max(
            (self._layer_max_parts(lyr) for lyr in (
                self.old_table_combo.currentLayer(),
                self.new_table_combo.currentLayer())), default=1)
        multipart = most_parts > 1

        self.geom_mode_combo.blockSignals(True)
        self.geom_mode_combo.clear()
        if multipart:
            self._add_mode_item("Singlepart", "parts")
            self._add_mode_item("Multipart", "whole")
            self.geom_mode_combo.setCurrentIndex(
                max(self.geom_mode_combo.findData(self._mode_choice), 0))
            self.geom_mode_combo.setItemData(
                0, "Split every feature into its parts and compare each "
                   "part separately", _ROLE_TOOLTIP)
            self.geom_mode_combo.setItemData(
                1, "Compare each feature as one whole shape", _ROLE_TOOLTIP)
            tip = (f"Multipart input detected (up to {most_parts:,} parts "
                   f"in one feature).\nChoose how it is compared.\n\n"
                   + MODE_TIP)
        else:
            self._add_mode_item("Singlepart", "whole")
            tip = ("Every feature has a single part, so there is nothing "
                   "to split.\n\n" + MODE_TIP)
        self.geom_mode_combo.blockSignals(False)
        self.geom_mode_combo.setEnabled(on and multipart)
        self.geom_mode_label.setEnabled(on and multipart)
        self.geom_mode_combo.setToolTip(tip)
        self.geom_mode_label.setToolTip(tip)

    def _on_mode_activated(self, _index):
        self._mode_choice = self.geom_mode_combo.currentData() or "parts"

    def _geom_by_part(self):
        return (self.geom_compare_cb.isChecked()
                and self.geom_mode_combo.isEnabled()
                and self.geom_mode_combo.currentData() == "parts")

    # ------------------------------------------------------------------
    # Review-points output location
    # ------------------------------------------------------------------
    def _set_review_tooltip(self):
        if self._review_path:
            tip = f"Review points are saved to:\n{self._review_path}"
        else:
            tip = (
                "A review point is created at every place that changed so "
                "you can\nvisit it.\n\n"
                "Temporary Points: a scratch layer that disappears when QGIS "
                "is closed.\nSave as shapefile / geopackage: keep the points "
                "in a file you choose.")
        self.review_out_label.setToolTip(tip)
        self.review_out_combo.setToolTip(tip)

    def _on_review_out_activated(self, index):
        fmt = self.review_out_combo.itemData(index)
        if fmt == "temp":
            self._review_path, self._review_fmt = "", "temp"
            self._set_review_tooltip()
            return

        ext = ".shp" if fmt == "shp" else ".gpkg"
        flt = ("Shapefile (*.shp)" if fmt == "shp"
               else "GeoPackage (*.gpkg)")
        start = (os.path.splitext(self._review_path)[0] + ext
                 if self._review_path else f"review_points{ext}")
        path, _sel = QFileDialog.getSaveFileName(
            self, "Save review points", start, flt)
        if not path:
            # cancelled: go back to the previous choice
            self.review_out_combo.setCurrentIndex(
                max(self.review_out_combo.findData(self._review_fmt), 0))
            return
        if not path.lower().endswith(ext):
            path = os.path.splitext(path)[0] + ext
        self._review_path, self._review_fmt = path, fmt
        self._set_review_tooltip()

    @staticmethod
    def _describe_layer(layer):
        if layer is None:
            return ""
        try:
            kind = (QgsWkbTypes.displayString(layer.wkbType())
                    if layer.isSpatial() else "Table")
        except Exception:
            kind = "Layer"
        count = layer.featureCount()
        count_txt = f"{count:,} features" if count >= 0 else "features: ?"
        return f"{kind}  ·  {count_txt}  ·  {layer.fields().count()} fields"

    def _update_layer_info(self):
        self.old_info.setText(
            self._describe_layer(self.old_table_combo.currentLayer()))
        self.new_info.setText(
            self._describe_layer(self.new_table_combo.currentLayer()))

    def _update_join_fields(self):
        previous = self.join_field_combo.currentText()
        self.join_field_combo.blockSignals(True)
        self.join_field_combo.clear()
        old_layer = self.old_table_combo.currentLayer()
        new_layer = self.new_table_combo.currentLayer()
        if old_layer and new_layer:
            new_names = {f.name() for f in new_layer.fields()}
            for fld in old_layer.fields():
                if fld.name() in new_names:
                    self.join_field_combo.addItem(
                        _field_icon(fld, self.theme), fld.name())
                    self.join_field_combo.setItemData(
                        self.join_field_combo.count() - 1,
                        _field_tooltip(fld), _ROLE_TOOLTIP)
            idx = self.join_field_combo.findText(previous)
            if idx >= 0:
                self.join_field_combo.setCurrentIndex(idx)
        self.join_field_combo.blockSignals(False)

    def _update_ready_state(self, *_args):
        old_layer = self.old_table_combo.currentLayer()
        new_layer = self.new_table_combo.currentLayer()
        ready = bool(
            old_layer and new_layer and self.join_field_combo.currentText())
        self.validate_btn.setEnabled(ready)
        self.compare_button.setEnabled(ready)
        self.column_filter_btn.setEnabled(bool(old_layer))

        if self.progress_bar.isVisible():
            return
        if not old_layer or not new_layer:
            self._set_status("Ready — select OLD and NEW layers to begin")
        elif old_layer.id() == new_layer.id():
            self._set_status(
                "OLD and NEW are the same layer — every feature will "
                "show as Unchanged", "warning")
        elif not ready:
            self._set_status(
                "The two layers share no field to use as a Join field",
                "warning")
        elif self._has_results:
            self._set_status(
                "Layer selection changed — press Compare to refresh "
                "the results")
        else:
            self._set_status("Ready — press Compare OLD → NEW")

    # ------------------------------------------------------------------
    # Geometry options
    # ------------------------------------------------------------------

    def _on_geom_toggle(self, enabled):
        """Enable / disable sub-checkboxes when master CB is toggled."""
        if enabled:
            self._update_geom_options()
        else:
            for cb in self._geom_checkboxes():
                cb.setEnabled(False)
        self._update_geom_mode()

    def _update_geom_options(self):
        """Tick the geometry checks that suit the OLD layer's geometry type.
        They are active only while 'Compare geometry' is ticked; otherwise
        they stay greyed but ticked, showing what would be compared."""
        on = self.geom_compare_cb.isChecked()
        old_layer = self.old_table_combo.currentLayer()
        if not old_layer:
            if not on:
                for cb in self._geom_checkboxes():
                    cb.setEnabled(False)
            return

        gt = old_layer.geometryType()
        is_line = (gt == _GEOM_LINE)
        is_polygon = (gt == _GEOM_POLYGON)

        for cb, applies in ((self.geom_length_cb, is_line or is_polygon),
                            (self.geom_area_cb, is_polygon),
                            (self.geom_centroid_cb, True),
                            (self.geom_vertices_cb, is_line or is_polygon)):
            cb.setEnabled(on and applies)
            cb.setChecked(applies)

    def _get_geom_field_names(self):
        """Return list of virtual geometry field names currently enabled."""
        names = []
        if not self.geom_compare_cb.isChecked():
            return names
        if self.geom_length_cb.isChecked() and self.geom_length_cb.isEnabled():
            names.append('_geom_length')
        if self.geom_area_cb.isChecked() and self.geom_area_cb.isEnabled():
            names.append('_geom_area')
        if self.geom_centroid_cb.isChecked() and self.geom_centroid_cb.isEnabled():
            names.extend(['_centroid_x', '_centroid_y'])
        if self.geom_vertices_cb.isChecked() and self.geom_vertices_cb.isEnabled():
            names.append('_vertex_count')
        return names

    def _add_geom_properties(self, data, geom):
        """Populate geometry virtual fields in *data* dict from *geom*."""
        if geom is None or geom.isNull() or geom.isEmpty():
            for fn in self._get_geom_field_names():
                data[fn] = None
            return

        if self.geom_length_cb.isChecked() and self.geom_length_cb.isEnabled():
            data['_geom_length'] = round(geom.length(), 6)

        if self.geom_area_cb.isChecked() and self.geom_area_cb.isEnabled():
            data['_geom_area'] = round(geom.area(), 6)

        if self.geom_centroid_cb.isChecked() and self.geom_centroid_cb.isEnabled():
            centroid_pt = geom.centroid().asPoint()
            data['_centroid_x'] = round(centroid_pt.x(), 6)
            data['_centroid_y'] = round(centroid_pt.y(), 6)

        if self.geom_vertices_cb.isChecked() and self.geom_vertices_cb.isEnabled():
            data['_vertex_count'] = sum(1 for _ in geom.vertices())

    # ====================================================================
    # PRE-FLIGHT VALIDATION
    # ====================================================================

    def _fit_validation(self):
        """Keep the report readable when the results table is full (up to
        130 px, then it scrolls) instead of squeezed to a sliver."""
        scroll, label = self.validation_scroll, self.validation_label
        width = max(scroll.viewport().width(), self.validation_frame.width() - 24, 200)
        needed = (label.heightForWidth(width) if label.hasHeightForWidth()
                  else label.sizeHint().height())
        scroll.setMinimumHeight(min(scroll.maximumHeight(),
                                    max(needed, label.fontMetrics().height()) + 2))

    def validate_data(self):
        """Run pre-flight checks and display results inline."""
        old_layer = self.old_table_combo.currentLayer()
        new_layer = self.new_table_combo.currentLayer()
        if not old_layer or not new_layer:
            QMessageBox.warning(
                self, "Warning",
                "Please select both OLD and NEW layers first!"
            )
            return

        join_field = self.join_field_combo.currentText()
        if not join_field:
            QMessageBox.warning(
                self, "Warning", "Please select a Join Field first!"
            )
            return

        ok_msgs, warn_msgs, err_msgs = [], [], []

        def collect_ids(layer):
            ids, nulls = [], 0
            for feat in layer.getFeatures():
                v = feat[join_field]
                if v is None or v == NULL or str(v).strip() == '':
                    nulls += 1
                else:
                    ids.append(v)
            return ids, nulls

        old_ids, old_nulls = collect_ids(old_layer)
        new_ids, new_nulls = collect_ids(new_layer)

        for label, ids in (("OLD", old_ids), ("NEW", new_ids)):
            dupes = {k: c for k, c in Counter(ids).items() if c > 1}
            if dupes:
                sample = ", ".join(
                    f"{k} (×{v})" for k, v in list(dupes.items())[:5]
                )
                if len(dupes) > 5:
                    sample += f", … +{len(dupes) - 5} more"
                err_msgs.append(
                    f"{label} has <b>{len(dupes)}</b> duplicate ID(s): "
                    f"{sample}"
                )
            else:
                ok_msgs.append(
                    f"{label}: No duplicate IDs ({len(ids)} unique records)"
                )

        for label, nulls in (("OLD", old_nulls), ("NEW", new_nulls)):
            if nulls:
                warn_msgs.append(
                    f"{label}: <b>{nulls}</b> record(s) with NULL / empty "
                    f"join field — these will be skipped"
                )

        old_fields = {f.name() for f in old_layer.fields()}
        new_fields = {f.name() for f in new_layer.fields()}
        only_old = old_fields - new_fields
        only_new = new_fields - old_fields

        ok_msgs.append(f"Common fields: <b>{len(old_fields & new_fields)}</b>")
        if only_old:
            warn_msgs.append(
                f"Fields only in OLD ({len(only_old)}): "
                f"{', '.join(sorted(only_old))}"
            )
        if only_new:
            warn_msgs.append(
                f"Fields only in NEW ({len(only_new)}): "
                f"{', '.join(sorted(only_new))}"
            )

        old_geom = QgsWkbTypes.displayString(old_layer.wkbType())
        new_geom = QgsWkbTypes.displayString(new_layer.wkbType())
        if old_layer.geometryType() == new_layer.geometryType():
            ok_msgs.append(f"Geometry types match: {old_geom}")
        else:
            warn_msgs.append(
                f"Geometry types differ: OLD = {old_geom}, NEW = {new_geom}"
            )

        t = self.theme
        html_parts = (
            [f'<span style="color:{t.success};">✓ {m}</span>'
             for m in ok_msgs]
            + [f'<span style="color:{t.warning};">⚠ {m}</span>'
               for m in warn_msgs]
            + [f'<span style="color:{t.danger};">✗ {m}</span>'
               for m in err_msgs]
        )
        self.validation_label.setText("<br>".join(html_parts))
        self.validation_frame.setVisible(True)
        self._fit_validation()
        self.clear_validation_btn.setEnabled(True)

        if err_msgs:
            self._set_status(
                f"Validation found {len(err_msgs)} error(s) — fix the "
                f"Join field or data before comparing", "error")
        elif warn_msgs:
            self._set_status(
                f"Validation passed with {len(warn_msgs)} warning(s)",
                "warning")
        else:
            self._set_status("Validation passed — no issues found",
                             "success")

    # ====================================================================
    # FILTERING
    # ====================================================================

    def _cell_values(self, row, col):
        """Searchable texts of a cell (both old and new for changed cells)."""
        item = self.results_table.item(row, col)
        if not item:
            return ()
        pair = item.data(_ROLE_USER) if col >= COL_FIRST_FIELD else None
        return pair if pair else (item.text(),)

    def _filter_columns(self):
        col = self.filter_field_combo.currentData()
        if col is None or col < 0:
            return range(COL_FIRST_FIELD, self.results_table.columnCount())
        return (col,)

    def _refresh_value_suggestions(self, *_args):
        """Rebuild the auto-complete list from the chosen field's values."""
        if not self._has_results:
            self.value_model.setStringList([])
            return
        values = set()
        for col in self._filter_columns():
            for row in range(self.results_table.rowCount()):
                for v in self._cell_values(row, col):
                    if v:
                        values.add(v)
                if len(values) >= MAX_SUGGESTIONS:
                    break
            if len(values) >= MAX_SUGGESTIONS:
                break
        self.value_model.setStringList(sorted(values, key=str.lower))

    def _on_filter_field_changed(self, *_args):
        # the box is narrow: its tooltip names the whole field
        self.filter_field_combo.setToolTip(
            f"Filter field: {self.filter_field_combo.currentText()}\n"
            "Choose the field the filter text is matched against")
        self._refresh_value_suggestions()
        self.apply_filters()

    def apply_filters(self, *_args):
        tbl = self.results_table
        if not self._has_results:
            return

        show = {s: b.isChecked() for s, b in self.filter_buttons.items()}
        needle = self.search_edit.text().strip().lower()
        cols = list(self._filter_columns())

        visible = 0
        for row in range(tbl.rowCount()):
            item = tbl.item(row, COL_STATUS)
            if not item:
                continue
            vis = show.get(item.text(), True)
            if vis and needle:
                vis = any(
                    needle in v.lower()
                    for c in cols for v in self._cell_values(row, c)
                )
            tbl.setRowHidden(row, not vis)
            if vis:
                visible += 1

        self._update_dynamic_row_numbers()
        self.rows_label.setText(f"Showing {visible:,} of {tbl.rowCount():,} rows")
        self._update_review_tally()

    # ====================================================================
    # VALUE FORMATTING & COMPARISON
    # ====================================================================

    @staticmethod
    def format_value(value):
        if value is None or value == NULL:
            return ""
        if isinstance(value, QDate):
            return value.toString("yyyy-MM-dd")
        if isinstance(value, QDateTime):
            return value.toString("yyyy-MM-dd hh:mm:ss")
        return str(value)

    @staticmethod
    def values_equal(val1, val2):
        if val1 is None or val1 == NULL:
            val1 = None
        if val2 is None or val2 == NULL:
            val2 = None

        if val1 is None and val2 is None:
            return True
        if val1 is None or val2 is None:
            return False

        s1, s2 = str(val1).strip(), str(val2).strip()
        try:
            return math.isclose(float(s1), float(s2), abs_tol=1e-10)
        except (ValueError, TypeError):
            return s1 == s2

    # ====================================================================
    # REVIEW (ACCEPT / REJECT)
    # ====================================================================

    def _style_decision_item(self, item, decision, status, progress=None):
        """Decision cell: Accepted / Rejected / Mixed, or Pending (with
        'decided / changed values' while a row is partly reviewed)."""
        t = self.theme
        item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
        item.setForeground(QBrush(QColor(t.row_text)))
        if decision == "Accepted":
            item.setText("✓ Accepted")
            item.setBackground(QBrush(t.accepted_bg))
        elif decision == "Rejected":
            item.setText("✗ Rejected")
            item.setBackground(QBrush(t.rejected_bg))
        elif decision == "Mixed":
            item.setText("◐ Mixed")
            item.setBackground(QBrush(t.bg(status)))
        else:
            item.setText(f"Pending {progress[0]}/{progress[1]}" if progress
                         else "Pending")
            item.setBackground(QBrush(t.bg(status)))

    # a changed value (pink cell with an OLD / NEW pair) can be decided alone
    @staticmethod
    def _reviewable(item):
        if item is None or item.column() < COL_FIRST_FIELD:
            return False
        pair = item.data(_ROLE_USER)
        return (item.data(_ROLE_KIND) == "changed"
                and isinstance(pair, (tuple, list)) and len(pair) == 2)

    def _row_values(self, row):
        tbl = self.results_table
        items = (tbl.item(row, c) for c in range(COL_FIRST_FIELD, tbl.columnCount()))
        return [it for it in items if self._reviewable(it)]

    def _set_value_decision(self, item, decision):
        """Accept (keep NEW), reject (keep OLD) or reset one changed value."""
        t = self.theme
        tbl = self.results_table
        fid = tbl.item(item.row(), COL_STATUS).data(_ROLE_USER)
        header = tbl.horizontalHeaderItem(item.column())
        name = header.text() if header else str(item.column())
        old_v, new_v = item.data(_ROLE_USER)
        if decision == "Accepted":
            self.cell_decisions[(fid, name)] = decision
            item.setText(new_v)
            item.setBackground(QBrush(t.bg("AcceptedValue")))
            note = "\nAccepted: the NEW value is kept"
        elif decision == "Rejected":
            self.cell_decisions[(fid, name)] = decision
            item.setText(old_v)
            item.setBackground(QBrush(t.bg("RejectedValue")))
            note = "\nRejected: the OLD value is kept"
        else:
            self.cell_decisions.pop((fid, name), None)
            item.setText(f"{old_v} → {new_v}")
            item.setBackground(QBrush(t.bg("ChangedField")))
            note = "\nRight-click to accept or reject"
        item.setData(_ROLE_DECISION, decision)
        item.setToolTip(f"{name}\nOLD: {old_v or '(empty)'}\n"
                        f"NEW: {new_v or '(empty)'}{note}")

    def _refresh_row_decision(self, row):
        """A row with changed values takes its decision from them."""
        tbl = self.results_table
        st = tbl.item(row, COL_STATUS)
        values = self._row_values(row)
        if not values:
            return
        fid = st.data(_ROLE_USER)
        states = [it.data(_ROLE_DECISION) for it in values]
        decided = sum(1 for s in states if s)
        if decided < len(states):
            self.decisions.pop(fid, None)
            decision = None
        elif all(s == "Accepted" for s in states):
            decision = "Accepted"
        elif all(s == "Rejected" for s in states):
            decision = "Rejected"
        else:
            decision = "Mixed"
        if decision:
            self.decisions[fid] = decision
        self._style_decision_item(
            tbl.item(row, COL_DECISION), decision, st.text(),
            (decided, len(states)) if decided and not decision else None)

    def _target_rows(self):
        tbl = self.results_table
        if self.review_scope.currentIndex() == 0:
            return sorted({it.row() for it in tbl.selectedItems()})
        return [r for r in range(tbl.rowCount()) if not tbl.isRowHidden(r)]

    def _apply_decision(self, decision):
        """Accept / Reject / Reset buttons (rows chosen by the scope box)."""
        rows = self._target_rows()
        if not rows:
            self._set_status(
                "Select one or more rows first (or switch to "
                "'All visible rows')", "warning")
            return
        self._decide_rows(rows, decision)

    def _decide_rows(self, rows, decision):
        """Decide whole rows: every changed value of a Modified row, or an
        Added row as a whole."""
        tbl = self.results_table
        changed = 0
        sorting = tbl.isSortingEnabled()
        tbl.setSortingEnabled(False)
        items = [(tbl.item(r, COL_STATUS), tbl.item(r, COL_DECISION),
                  self._row_values(r)) for r in rows]
        for st, dec_item, values in items:
            if not st or st.text() not in REVIEWABLE or not dec_item:
                continue
            if values:
                for item in values:
                    self._set_value_decision(item, decision)
                self._refresh_row_decision(st.row())
            else:
                fid = st.data(_ROLE_USER)
                if decision is None:
                    self.decisions.pop(fid, None)
                else:
                    self.decisions[fid] = decision
                self._style_decision_item(dec_item, decision, st.text())
            changed += 1
        tbl.setSortingEnabled(sorting)

        self._update_review_tally()
        if changed:
            verb = {"Accepted": "Accepted", "Rejected": "Rejected",
                    None: "Reset to Pending"}[decision]
            self._set_status(f"{verb}: {changed:,} row(s)", "success")
        else:
            self._set_status(
                "Only Added and Modified rows can be reviewed", "warning")

    def _decide_values(self, decision):
        """Accept / reject / reset the selected changed values only."""
        tbl = self.results_table
        values = [it for it in tbl.selectedItems() if self._reviewable(it)]
        if not values:
            self._set_status(
                "Select one or more highlighted (changed) values first",
                "warning")
            return
        sorting = tbl.isSortingEnabled()
        tbl.setSortingEnabled(False)
        statuses = [tbl.item(it.row(), COL_STATUS) for it in values]
        for item in values:
            self._set_value_decision(item, decision)
        for st in {id(s): s for s in statuses}.values():
            self._refresh_row_decision(st.row())
        tbl.setSortingEnabled(sorting)
        self._update_review_tally()
        verb = {"Accepted": "Accepted", "Rejected": "Rejected",
                None: "Reset to Pending"}[decision]
        self._set_status(f"{verb}: {len(values):,} value(s)", "success")

    def _update_review_tally(self):
        if not self._has_results:
            self.review_tally.setText("")
            return
        counts = Counter(self.decisions.values())
        acc, rej, mixed = counts["Accepted"], counts["Rejected"], counts["Mixed"]
        pend = max(self._reviewable_total - acc - rej - mixed, 0)
        t = self.theme
        text = (f"<span style='color:{t.accent_of('AcceptedValue').name()};'>"
                f"Accepted {acc:,}</span>"
                f"  ·  <span style='color:{t.accent_of('RejectedValue').name()};'>"
                f"Rejected {rej:,}</span>")
        if mixed:
            text += f"  ·  Mixed {mixed:,}"
        self.review_tally.setText("Rows: " + text + f"  ·  Pending {pend:,}")
        values = Counter(self.cell_decisions.values())
        pending = max(self._value_total - values["Accepted"] - values["Rejected"], 0)
        self.changed_field_legend.setText(f"Changed · {pending:,}")
        self.accepted_legend.setText(f"Accepted · {values['Accepted']:,}")
        self.rejected_legend.setText(f"Rejected · {values['Rejected']:,}")

    # ------------------------------------------------------------------
    # Right-click menu on the results
    # ------------------------------------------------------------------

    def _on_table_menu(self, pos):
        tbl = self.results_table
        item = tbl.itemAt(pos)
        if item is None:
            return
        if not item.isSelected():
            tbl.clearSelection()
            tbl.setCurrentItem(item)
        selected = tbl.selectedItems()
        values = [it for it in selected if self._reviewable(it)]
        rows = sorted({it.row() for it in selected})
        review_rows = [r for r in rows
                       if tbl.item(r, COL_STATUS).text() in REVIEWABLE]

        def label(verb, n, noun):
            return f"{verb} {noun}" if n == 1 else f"{verb} {n:,} {noun}s"

        menu = QMenu(tbl)
        choices = (("Accept", icons.SUCCESS, "Accepted"),
                   ("Reject", icons.DELETE, "Rejected"),
                   ("Reset", icons.UNDO, None))
        if values:
            for verb, icon, decision in choices:
                menu.addAction(icons.native(icon),
                               label(verb, len(values), "value")).triggered.connect(
                    lambda _c=False, d=decision: self._decide_values(d))
            menu.addSeparator()
        if review_rows:
            for verb, icon, decision in choices:
                menu.addAction(icons.native(icon),
                               label(verb, len(review_rows), "row")).triggered.connect(
                    lambda _c=False, d=decision, r=review_rows: self._decide_rows(r, d))
            menu.addSeparator()
        zoom = menu.addAction(icons.native(icons.ZOOM_SELECTED),
                              label("Zoom to", len(rows), "feature"))
        zoom.triggered.connect(self.zoom_to_selection)
        if self._geom_diffs:
            menu.addAction(icons.native(icons.GEOM_DIFF), "Show changes on map"
                           ).triggered.connect(self.show_changes_on_map)
        menu.exec(tbl.viewport().mapToGlobal(pos))

    # ------------------------------------------------------------------
    # Zoom to / select the features of the selected rows
    # ------------------------------------------------------------------

    def zoom_to_selection(self):
        tbl = self.results_table
        rows = sorted({it.row() for it in tbl.selectedItems()})
        if not rows:
            self._set_status("Select one or more rows first", "warning")
            return
        project = QgsProject.instance()
        old_layer = project.mapLayer(self._run_layer_ids[0] or "")
        new_layer = project.mapLayer(self._run_layer_ids[1] or "")
        if old_layer is None or new_layer is None:
            self._set_status("The compared layers are no longer in the project "
                             "— press Compare again", "warning")
            return
        old_map, new_map = self._fid_maps["old"], self._fid_maps["new"]
        old_ids, new_ids, old_only = [], [], []
        for r in rows:
            key = tbl.item(r, COL_STATUS).data(_ROLE_USER)
            if key in new_map:
                new_ids.append(new_map[key])
            if key in old_map:
                old_ids.append(old_map[key])
                if key not in new_map:
                    old_only.append(old_map[key])
        if old_layer.id() == new_layer.id():
            new_layer.selectByIds(sorted(set(old_ids) | set(new_ids)), SELECT_SET)
        else:
            old_layer.selectByIds(old_ids, SELECT_SET)
            new_layer.selectByIds(new_ids, SELECT_SET)

        canvas = self.iface.mapCanvas()
        dest = canvas.mapSettings().destinationCrs()
        extent = QgsRectangle()
        extent.setNull()
        flashes = []
        for layer, ids in ((new_layer, new_ids), (old_layer, old_only)):
            if not ids or not layer.isSpatial():
                continue
            request = QgsFeatureRequest().setFilterFids(sorted(ids))
            geoms = [QgsGeometry(f.geometry()) for f in layer.getFeatures(request)
                     if _geom_present(f.geometry())]
            if not geoms:
                continue
            to_canvas = QgsCoordinateTransform(layer.crs(), dest, project)
            for g in geoms:
                try:
                    extent.combineExtentWith(
                        to_canvas.transformBoundingBox(g.boundingBox()))
                except Exception:
                    extent.combineExtentWith(g.boundingBox())
            flashes.append((geoms[:200], layer.crs()))

        n = len(set(old_ids) | set(new_ids)) if old_layer.id() == new_layer.id() \
            else len(rows)
        if not flashes:
            self._set_status(f"Selected the features of {n:,} row(s) "
                             "— there is no geometry to zoom to", "success")
            return
        if extent.width() == 0 or extent.height() == 0:
            pad = max(extent.width(), extent.height(), 1e-4) * 0.5
            extent.grow(pad)
        extent.scale(1.25)
        canvas.setExtent(extent)
        canvas.refresh()
        for geoms, crs in flashes:
            try:
                canvas.flashGeometries(geoms, crs)
            except (RuntimeError, TypeError):
                pass                     # flashing is only a visual hint
        self._set_status(
            f"Selected and zoomed to {n:,} row(s) — {len(new_ids):,} in NEW, "
            f"{len(old_ids):,} in OLD", "success")

    # ------------------------------------------------------------------
    # Help and validation
    # ------------------------------------------------------------------

    def show_help(self):
        from .help_dialog import HelpDialog
        if self._help_dialog is None:
            self._help_dialog = HelpDialog(self, topic="compare")
        self._help_dialog.show()
        self._help_dialog.raise_()
        self._help_dialog.activateWindow()

    def clear_validation(self):
        self.validation_label.setText("")
        self.validation_frame.setVisible(False)
        self.clear_validation_btn.setEnabled(False)
        self._set_status("Validation results cleared")

    # ====================================================================
    # SHOW CHANGES ON THE MAP
    # ====================================================================

    _MAP_STYLE = {
        "added": (f"{PLUGIN_NAME} – Added (NEW − OLD)", "#2E9E4F"),
        "removed": (f"{PLUGIN_NAME} – Removed (OLD − NEW)", "#E03C31"),
    }

    def _rows_for_map(self):
        tbl = self.results_table
        selected = sorted({it.row() for it in tbl.selectedItems()})
        rows = selected or [r for r in range(tbl.rowCount())
                            if not tbl.isRowHidden(r)]
        fids = []
        for r in rows:
            item = tbl.item(r, COL_STATUS)
            if item and item.text() != "Unchanged":
                fids.append(item.data(_ROLE_USER))
        return fids, bool(selected)

    def _build_diff_layer(self, kind, fids):
        """Memory layer holding the added / removed geometry of *fids*."""
        name, colour = self._MAP_STYLE[kind]
        gtype = self._geom_gtype
        wkb = {_GEOM_POLYGON: "MultiPolygon", _GEOM_LINE: "MultiLineString",
               _GEOM_POINT: "MultiPoint"}[gtype]
        layer = QgsVectorLayer(
            f"{wkb}?field=id:string&field=change:string"
            f"&field=measure:double", name, "memory")
        layer.setCrs(self._geom_crs)

        feats = []
        for fid in fids:
            gd = self._geom_diffs.get(fid)
            geom = gd.get(kind) if gd else None
            if not _geom_present(geom):
                continue
            geom = QgsGeometry(geom)
            geom.convertToMultiType()
            feat = QgsFeature(layer.fields())
            feat.setGeometry(geom)
            feat.setAttributes([
                str(fid), kind,
                gd["added_m"] if kind == "added" else gd["removed_m"]])
            feats.append(feat)
        if not feats:
            return None
        layer.dataProvider().addFeatures(feats)
        layer.updateExtents()

        symbol = QgsSymbol.defaultSymbol(layer.geometryType())
        symbol.setColor(QColor(colour))
        if gtype == _GEOM_LINE:
            symbol.setWidth(0.9)
        elif gtype == _GEOM_POINT:
            symbol.setSize(3.5)
        layer.setRenderer(QgsSingleSymbolRenderer(symbol))
        layer.setOpacity(0.8)
        return layer

    def _on_row_double_click(self, row, _col):
        """Double-click a changed row: zoom to where it changed."""
        item = self.results_table.item(row, COL_STATUS)
        if not item or item.text() == "Unchanged" or not self._geom_diffs:
            return
        self.results_table.clearSelection()
        self.results_table.selectRow(row)
        self.show_changes_on_map()

    def show_changes_on_map(self):
        """Add the added / removed geometry of the chosen rows to the map and
        zoom to it, so changes inside a huge multipart feature can be found."""
        if not self._geom_diffs:
            self._set_status(
                "Tick 'Compare geometry' and run Compare to use this",
                "warning")
            return
        fids, from_selection = self._rows_for_map()
        fids = [f for f in fids if self._geom_diffs.get(f) and (
            self._geom_diffs[f]["added"] or self._geom_diffs[f]["removed"])]
        if not fids:
            self._set_status(
                "No geometry changes in the "
                f"{'selected' if from_selection else 'visible'} rows",
                "warning")
            return

        project = QgsProject.instance()
        for old_id in self._map_layer_ids:
            if project.mapLayer(old_id):
                project.removeMapLayer(old_id)
        self._map_layer_ids = []

        layers = []
        for kind in ("added", "removed"):
            layer = self._build_diff_layer(kind, fids)
            if layer is not None:
                layers.append(layer)

        root = project.layerTreeRoot()
        group = root.findGroup(MAP_GROUP_NAME) or root.insertGroup(
            0, MAP_GROUP_NAME)
        for layer in layers:
            project.addMapLayer(layer, False)
            group.addLayer(layer)
            self._map_layer_ids.append(layer.id())

        extent = QgsRectangle()
        extent.setNull()
        for layer in layers:
            extent.combineExtentWith(layer.extent())
        self._zoom_map_to(extent, layers)

        n_added = sum(1 for f in fids if self._geom_diffs[f]["added"])
        n_removed = sum(1 for f in fids if self._geom_diffs[f]["removed"])
        self._set_status(
            f"Map: {n_added:,} feature(s) with added geometry (green), "
            f"{n_removed:,} with removed geometry (red) — "
            f"{'selected rows' if from_selection else 'visible rows'}",
            "success")

    def _zoom_map_to(self, extent, layers):
        if extent.isNull():
            return
        canvas = self.iface.mapCanvas()
        try:
            to_canvas = QgsCoordinateTransform(
                self._geom_crs, canvas.mapSettings().destinationCrs(),
                QgsProject.instance())
            extent = to_canvas.transformBoundingBox(extent)
        except QgsCsException:
            pass                         # keep the extent in the layer CRS
        if extent.width() == 0 or extent.height() == 0:
            pad = max(extent.width(), extent.height(), 1e-4) * 0.5
            extent.grow(pad)
        extent.scale(1.25)
        canvas.setExtent(extent)
        canvas.refresh()
        try:
            geoms = [f.geometry() for lyr in layers
                     for f in lyr.getFeatures()][:200]
            canvas.flashGeometries(geoms, self._geom_crs)
        except (RuntimeError, TypeError):
            pass                         # flashing is only a visual hint

    # ====================================================================
    # EXPORT
    # ====================================================================

    # Excel always uses the light palette (white sheet) so it prints well
    _XL_STATUS = Theme.LIGHT_STATUS
    _XL_HEADER = "E4E4E4"
    _XL_ACCEPTED = Theme.LIGHT_STATUS["AcceptedValue"][0].lstrip("#")
    _XL_REJECTED = Theme.LIGHT_STATUS["RejectedValue"][0].lstrip("#")

    def export_xlsx(self):
        """Export exactly what the table shows (visible rows, same order),
        with the same cell colours, to an Excel workbook."""
        tbl = self.results_table
        if tbl.rowCount() == 0:
            QMessageBox.warning(self, "Warning", "No data to export!")
            return
        if not HAS_OPENPYXL:
            QMessageBox.warning(
                self, "Excel export unavailable",
                "The 'openpyxl' Python package is not installed.\n"
                "Use the CSV export instead.")
            return

        filename, _ = QFileDialog.getSaveFileName(
            self, "Export to Excel", "comparison_results.xlsx",
            "Excel files (*.xlsx)"
        )
        if not filename:
            return
        if not filename.lower().endswith(".xlsx"):
            filename += ".xlsx"

        try:
            exported = self._write_xlsx(filename)
            self._set_status(f"Excel exported: {filename}", "success")
            QMessageBox.information(
                self, "Export Complete",
                f"Excel exported to:\n{filename}\n{exported:,} row(s)."
            )
        except PermissionError:
            self._set_status("Export failed — file is open", "error")
            QMessageBox.critical(
                self, "Error",
                "Could not write the file. Is it open in Excel?")
        except Exception as e:
            self._set_status("Export failed", "error")
            QMessageBox.critical(
                self, "Error", f"Failed to export:\n{str(e)}")

    def _write_xlsx(self, filename):
        tbl = self.results_table
        ncols = tbl.columnCount()
        side = Side(style="thin", color="CCCCCC")
        border = Border(left=side, right=side, top=side, bottom=side)
        font = XlFont(name="Calibri", size=10, color="1B1B1B")
        bold = XlFont(name="Calibri", size=10, bold=True, color="1B1B1B")
        centre = Alignment(horizontal="center", vertical="center")
        left = Alignment(horizontal="left", vertical="center")

        def fill(hex_colour):
            return PatternFill(
                start_color=hex_colour, end_color=hex_colour,
                fill_type="solid")

        def hex_of(key):
            return self._XL_STATUS[key][0].lstrip("#")

        wb = Workbook()
        ws = wb.active
        ws.title = "Comparison"

        # ---- header: '#' + the table's own columns ----
        headers = ["#"]
        for c in range(ncols):
            hi = tbl.horizontalHeaderItem(c)
            headers.append(hi.text() if hi else f"Col_{c}")
        for ci, text in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=text)
            cell.font = bold
            cell.fill = fill(self._XL_HEADER)
            cell.alignment = centre
            cell.border = border

        # ---- rows: visible only, in the order currently displayed ----
        out_row = 1
        for r in range(tbl.rowCount()):
            if tbl.isRowHidden(r):
                continue
            out_row += 1
            status_item = tbl.item(r, COL_STATUS)
            status = status_item.text()
            fid = status_item.data(_ROLE_USER)
            decision = self.decisions.get(fid)

            num = ws.cell(row=out_row, column=1, value=out_row - 1)
            num.font = bold
            num.fill = fill(self._XL_HEADER)
            num.alignment = centre
            num.border = border

            for c in range(ncols):
                item = tbl.item(r, c)
                cell = ws.cell(
                    row=out_row, column=c + 2,
                    value=item.text() if item else "")
                cell.font = font
                cell.border = border
                cell.alignment = centre if c in (COL_STATUS, COL_DECISION) \
                    else left
                value_decision = item.data(_ROLE_DECISION) if item else None
                if c == COL_DECISION and decision == "Accepted":
                    cell.fill = fill(self._XL_ACCEPTED)
                elif c == COL_DECISION and decision == "Rejected":
                    cell.fill = fill(self._XL_REJECTED)
                elif c >= COL_FIRST_FIELD and value_decision == "Accepted":
                    cell.fill = fill(self._XL_ACCEPTED)
                elif c >= COL_FIRST_FIELD and value_decision == "Rejected":
                    cell.fill = fill(self._XL_REJECTED)
                elif (c >= COL_FIRST_FIELD and item
                        and item.data(_ROLE_KIND) == "changed"):
                    cell.fill = fill(hex_of("ChangedField"))
                else:
                    cell.fill = fill(hex_of(status))

        # ---- sizing / navigation ----
        for ci in range(1, len(headers) + 1):
            longest = max(
                len(str(ws.cell(row=ri, column=ci).value or ""))
                for ri in range(1, out_row + 1))
            header_len = len(headers[ci - 1]) + 6   # room for filter arrow
            ws.column_dimensions[get_column_letter(ci)].width = \
                min(max(longest + 3, header_len, 7), 50)
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions

        self._write_xlsx_summary(wb, out_row - 1, fill, border, font, bold)
        wb.save(filename)
        return out_row - 1

    def _write_xlsx_summary(self, wb, exported, fill, border, font, bold):
        info = self._compare_info
        ws = wb.create_sheet("Summary")
        counts = info.get("counts", {})

        ws.cell(row=1, column=1, value=PLUGIN_NAME).font = XlFont(
            name="Calibri", size=14, bold=True)
        meta = [
            ("Exported", QDateTime.currentDateTime().toString(
                "yyyy-MM-dd hh:mm")),
            ("OLD (baseline)", info.get("old_name", "")),
            ("NEW (updated)", info.get("new_name", "")),
            ("Join field", info.get("join_field", "")),
            ("Fields compared", ", ".join(info.get("checked_fields", []))),
            ("Geometry compared", ", ".join(info.get("geometry", [])) or "—"),
            ("Geometry tolerance",
             "—" if info.get("tolerance") is None
             else f"{info['tolerance']:g} (layer map units)"),
            ("Comparison method", (info.get("analysis") or "—")),
            ("Review points", info.get("review_path")
             or ("temporary layer" if info.get("review") else "—")),
            ("Rows exported",
             f"{exported:,} of {self.results_table.rowCount():,} "
             f"(filters applied)"),
        ]
        row = 3
        for label, value in meta:
            ws.cell(row=row, column=1, value=label).font = bold
            value_cell = ws.cell(row=row, column=2, value=value)
            value_cell.font = font
            value_cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=row, column=1).alignment = Alignment(vertical="top")
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Legend").font = bold
        row += 1
        legend = [
            ("Added", "Added", "Only in NEW"),
            ("Deleted", "Deleted", "Only in OLD"),
            ("Modified", "Modified", "In both, a compared value differs"),
            ("Unchanged", "Unchanged", "Identical in both"),
            ("ChangedField", "Changed field",
             "Cell shows  OLD value → NEW value  (blank = empty)"),
            ("AcceptedValue", "Accepted value",
             "Change accepted: the cell shows the NEW value"),
            ("RejectedValue", "Rejected value",
             "Change rejected: the cell shows the OLD value"),
        ]
        for key, label, desc in legend:
            swatch = ws.cell(row=row, column=1, value=label)
            swatch.font = font
            swatch.border = border
            swatch.fill = fill(self._XL_STATUS[key][0].lstrip("#"))
            swatch.alignment = Alignment(horizontal="center")
            if key in counts:
                ws.cell(row=row, column=2,
                        value=f"{counts[key]:,}  —  {desc}").font = font
            else:
                ws.cell(row=row, column=2, value=desc).font = font
            row += 1

        decided = Counter(self.decisions.values())
        for label, colour, n in (("Accepted", self._XL_ACCEPTED, decided["Accepted"]),
                                 ("Rejected", self._XL_REJECTED, decided["Rejected"]),
                                 ("Mixed", self._XL_HEADER, decided["Mixed"])):
            swatch = ws.cell(row=row, column=1, value=label)
            swatch.font = font
            swatch.border = border
            swatch.fill = fill(colour)
            swatch.alignment = Alignment(horizontal="center")
            ws.cell(row=row, column=2,
                    value=f"{n:,}  —  Decision column").font = font
            row += 1

        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 70

    def export_results(self):
        """Export the visible rows to CSV (plain values)."""
        tbl = self.results_table
        if tbl.rowCount() == 0:
            QMessageBox.warning(self, "Warning", "No data to export!")
            return

        filename, _ = QFileDialog.getSaveFileName(
            self, "Export Comparison Results",
            "comparison_results.csv", "CSV files (*.csv)"
        )
        if not filename:
            return

        try:
            cols = [COL_STATUS] + list(range(COL_FIRST_FIELD,
                                             tbl.columnCount()))
            with open(filename, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)

                headers = []
                for c in cols:
                    hi = tbl.horizontalHeaderItem(c)
                    headers.append(hi.text() if hi else f"Col_{c}")
                headers.append("Decision")
                writer.writerow(headers)

                for r in range(tbl.rowCount()):
                    if tbl.isRowHidden(r):
                        continue
                    fid = tbl.item(r, COL_STATUS).data(_ROLE_USER)
                    decision = self.decisions.get(fid, "Pending")

                    row_data = []
                    for c in cols:
                        item = tbl.item(r, c)
                        if not item:
                            row_data.append("")
                            continue
                        pair = item.data(_ROLE_USER) if c != COL_STATUS else None
                        if pair:
                            # a rejected value keeps OLD; accepted / pending NEW
                            old_v, new_v = pair
                            rejected = (item.data(_ROLE_DECISION) == "Rejected"
                                        or (decision == "Rejected"
                                            and not self._reviewable(item)))
                            row_data.append(old_v if rejected else new_v)
                        else:
                            row_data.append(item.text())
                    row_data.append(decision)
                    writer.writerow(row_data)

            self._set_status(f"Exported results to {filename}", "success")
            QMessageBox.information(
                self, "Success", f"Results exported to:\n{filename}"
            )
        except Exception as e:
            self._set_status("Export failed", "error")
            QMessageBox.critical(
                self, "Error", f"Failed to export:\n{str(e)}"
            )

    # ====================================================================
    # COLUMN SELECTION
    # ====================================================================

    def select_columns_to_check(self):
        old_layer = self.old_table_combo.currentLayer()
        if not old_layer:
            QMessageBox.warning(
                self, "Warning",
                "Please select the OLD Version layer first!"
            )
            return

        fields = list(old_layer.fields())
        if not fields:
            QMessageBox.warning(
                self, "Warning", "No fields found in selected layer!"
            )
            return

        if self.columns_to_check:
            checked = set(self.columns_to_check)
        else:
            checked = {f.name() for f in fields
                       if f.name().lower() not in AUTO_SKIP_FIELDS}

        entries = [(f.name(), _field_icon(f, self.theme), _field_tooltip(f),
                    f.name() in checked) for f in fields]
        dlg = W.FieldChecklistDialog(
            entries, self.metrics, self.theme, "Field(s) to Compare",
            "Choose the fields used to decide whether a feature is "
            "<b>Modified</b>. Unticked fields are ignored "
            "(e.g. fid, timestamps).", parent=self)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            self.columns_to_check = dlg.checked_names()
            if self._has_results:
                self.compare_tables()

    # ====================================================================
    # SORTING & ROW NUMBERING
    # ====================================================================

    def _on_column_sort(self):
        QTimer.singleShot(10, self._update_dynamic_row_numbers)

    def _update_dynamic_row_numbers(self):
        vis = 0
        for r in range(self.results_table.rowCount()):
            if not self.results_table.isRowHidden(r):
                vis += 1
                self.results_table.setVerticalHeaderItem(
                    r, QTableWidgetItem(str(vis))
                )

    # ====================================================================
    # MAIN COMPARISON LOGIC
    # ====================================================================

    def compare_tables(self):
        old_layer = self.old_table_combo.currentLayer()
        new_layer = self.new_table_combo.currentLayer()
        if not old_layer or not new_layer:
            QMessageBox.warning(
                self, "Warning",
                "Please select both OLD and NEW layers!"
            )
            return

        join_field = self.join_field_combo.currentText()
        if not join_field:
            QMessageBox.warning(
                self, "Warning", "No Join Field available!"
            )
            return

        self.progress_bar.setVisible(True)
        self.progress_bar.setValue(0)
        self.compare_button.setEnabled(False)
        self.validate_btn.setEnabled(False)
        self.compare_button.setText("Processing…")
        self._set_status("Comparing tables…", "busy")
        QApplication.processEvents()

        try:
            self.results_table.setRowCount(0)
            self.results_table.setColumnCount(0)
            self.decisions = {}
            self.cell_decisions = {}
            self._fid_maps = {"old": {}, "new": {}}
            self._run_layer_ids = (old_layer.id(), new_layer.id())

            new_field_set = {f.name() for f in new_layer.fields()}
            attr_fields = [
                f.name() for f in old_layer.fields()
                if f.name() in new_field_set
            ]

            if not self.columns_to_check:
                self.columns_to_check = [
                    f for f in attr_fields if f.lower() not in AUTO_SKIP_FIELDS
                ]

            geom_fields = self._get_geom_field_names()
            self._old_fields = old_layer.fields()

            # Exact (part-aware) geometry comparison
            gtype = old_layer.geometryType()
            geom_diff_on = (
                self.geom_compare_cb.isChecked()
                and gtype in (_GEOM_POINT, _GEOM_LINE, _GEOM_POLYGON)
                and new_layer.geometryType() == gtype
            )
            old_geoms, new_geoms = {}, {}
            to_new_crs = None
            if geom_diff_on and old_layer.crs() != new_layer.crs():
                to_new_crs = QgsCoordinateTransform(
                    old_layer.crs(), new_layer.crs(), QgsProject.instance())

            self.progress_bar.setValue(5)
            QApplication.processEvents()

            old_features = self._read_features(
                old_layer, join_field, attr_fields, geom_fields,
                old_geoms if geom_diff_on else None, to_new_crs,
                self._fid_maps["old"])
            self.progress_bar.setValue(25)
            QApplication.processEvents()

            new_features = self._read_features(
                new_layer, join_field, attr_fields, geom_fields,
                new_geoms if geom_diff_on else None, None,
                self._fid_maps["new"])
            self.progress_bar.setValue(45)
            QApplication.processEvents()

            geom_diffs = {}
            progress_base = 45
            by_part = self._geom_by_part()
            self._geom_by_part_run = bool(geom_diff_on and by_part)
            self._review_note = ""
            if geom_diff_on:
                geom_diffs = self._compute_geom_diffs(
                    old_geoms, new_geoms, gtype, by_part)
                progress_base = 70
            self._geom_diffs = geom_diffs
            self._geom_crs = new_layer.crs()
            self._geom_gtype = gtype if geom_diff_on else None
            del old_geoms, new_geoms

            if geom_diff_on:
                self._review_note = self._create_review_points(geom_diffs)
            else:
                self._remove_review_layer()

            self._display_results(
                old_features, new_features,
                attr_fields + geom_fields,
                list(self.columns_to_check) + geom_fields,
                old_layer.name(), new_layer.name(),
                geom_diffs, gtype, progress_base
            )
            self.progress_bar.setValue(100)

        except Exception as e:
            QMessageBox.critical(
                self, "Error", f"Comparison failed:\n{str(e)}"
            )
            self._set_status("Error during comparison", "error")
        finally:
            self.compare_button.setText("Compare OLD → NEW")
            # bound to the progress bar, so Qt drops the timer if the panel
            # is closed within the delay (a lambda would hit a deleted widget)
            QTimer.singleShot(1500, self.progress_bar.hide)
            self._update_ready_state()

    def _read_features(self, layer, join_field, attr_fields, geom_fields,
                       geom_store=None, transform=None, id_store=None):
        features = {}
        for feat in layer.getFeatures():
            key = feat[join_field]
            if key is None or key == NULL:
                continue
            if id_store is not None:
                id_store[key] = feat.id()
            data = {fn: feat[fn] for fn in attr_fields}
            if geom_fields:
                self._add_geom_properties(data, feat.geometry())
            if geom_store is not None:
                geom = QgsGeometry(feat.geometry())
                if transform is not None and _geom_present(geom):
                    geom.transform(transform)
                geom_store[key] = geom
            features[key] = data
        return features

    def _compute_geom_diffs(self, old_geoms, new_geoms, gtype, by_part):
        """Per-feature geometry analysis (whole feature or part by part)."""
        tol = self.geom_tol_spin.value()
        ids = list(set(old_geoms) | set(new_geoms))
        diffs = {}
        total = max(len(ids), 1)
        for i, fid in enumerate(ids):
            diffs[fid] = analyse_geometry(
                old_geoms.get(fid), new_geoms.get(fid), tol, gtype, by_part)
            if i % 5 == 0:
                self.progress_bar.setValue(45 + int(25 * (i + 1) / total))
                QApplication.processEvents()
        return diffs

    # ====================================================================
    # REVIEW POINTS (where did it change?)
    # ====================================================================

    _REVIEW_COLOURS = {"added": "#2E9E4F", "removed": "#E03C31",
                       "modified": "#F29B1D"}

    def _remove_review_layer(self):
        project = QgsProject.instance()
        if self._review_layer_id and project.mapLayer(self._review_layer_id):
            project.removeMapLayer(self._review_layer_id)
        self._review_layer_id = None

    def _style_review_layer(self, layer):
        cats = []
        for value, colour in self._REVIEW_COLOURS.items():
            symbol = QgsMarkerSymbol.createSimple({
                "name": "diamond", "color": colour,
                "outline_color": "#FFFFFF", "outline_width": "0.5",
                "size": "4"})
            cats.append(QgsRendererCategory(value, symbol, value))
        layer.setRenderer(QgsCategorizedSymbolRenderer("change", cats))

    def _write_review_file(self, layer, path):
        """Save *layer* as .gpkg / .shp; returns (ok, message)."""
        gpkg = path.lower().endswith(".gpkg")
        options = QgsVectorFileWriter.SaveVectorOptions()
        options.driverName = "GPKG" if gpkg else "ESRI Shapefile"
        options.fileEncoding = "UTF-8"
        if gpkg:
            options.layerName = "review_points"
            if os.path.exists(path):
                # re-running Compare replaces only this layer, so other
                # layers already stored in the same geopackage are kept
                action = getattr(QgsVectorFileWriter, "ActionOnExistingFile",
                                 QgsVectorFileWriter)
                options.actionOnExistingFile = action.CreateOrOverwriteLayer
        result = QgsVectorFileWriter.writeAsVectorFormatV3(
            layer, path, QgsProject.instance().transformContext(), options)
        return result[0] == QgsVectorFileWriter.WriterError.NoError, (result[1] or "")

    def _create_review_points(self, geom_diffs):
        """Build the review-point layer (temporary or saved) and add it to
        the project.  Returns a short note for the status bar."""
        self._remove_review_layer()
        records = []
        for fid, gd in geom_diffs.items():
            for rec in gd.get("points", ()):
                records.append((fid, rec))
        if not records:
            return "no changed geometry"

        unit = _geom_unit(self._geom_gtype)
        layer = QgsVectorLayer(
            "Point?" + "&".join(f"field={n}:{t}" for n, t in REVIEW_FIELDS),
            REVIEW_LAYER_NAME, "memory")
        layer.setCrs(self._geom_crs)
        feats = []
        for fid, rec in records:
            feat = QgsFeature(layer.fields())
            feat.setGeometry(rec["geom"])
            feat.setAttributes([
                str(fid), rec["change"], rec["level"],
                rec["part_old"], rec["part_new"], rec["measure"], unit])
            feats.append(feat)
        layer.dataProvider().addFeatures(feats)
        layer.updateExtents()

        note = f"{len(records):,} review point(s)"
        shown = layer
        self._review_saved = ""
        if self._review_path:
            ok, _msg = self._write_review_file(layer, self._review_path)
            saved = None
            if ok:
                uri = self._review_path
                if uri.lower().endswith(".gpkg"):
                    uri += "|layername=review_points"
                saved = QgsVectorLayer(uri, REVIEW_LAYER_NAME, "ogr")
            if saved is not None and saved.isValid():
                shown = saved
                self._review_saved = self._review_path
                note += f" saved to {os.path.basename(self._review_path)}"
            else:
                note += (f" — could not save to "
                         f"{os.path.basename(self._review_path)}, kept as a "
                         f"temporary layer")
        self._style_review_layer(shown)

        project = QgsProject.instance()
        root = project.layerTreeRoot()
        group = root.findGroup(MAP_GROUP_NAME) or root.insertGroup(
            0, MAP_GROUP_NAME)
        project.addMapLayer(shown, False)
        group.insertLayer(0, shown)
        self._review_layer_id = shown.id()
        return note

    # ------------------------------------------------------------------

    def _new_row_item(self, text, status):
        item = _SortItem(text)
        item.setBackground(QBrush(self.theme.bg(status)))
        item.setForeground(QBrush(QColor(self.theme.row_text)))
        return item

    @staticmethod
    def _geom_extra_names(gtype, by_part=False):
        unit = _geom_unit(gtype)
        names = ["_geom_status", f"_added_{unit}", f"_removed_{unit}",
                 "_changed_pct", "_parts"]
        if by_part:
            names.append("_parts_changed")
        return names

    def _geom_extra_cell(self, name, status, gd, unit):
        """(text, pink?, (old, new) pair or None) for a geometry-diff column."""
        if gd is None:
            return "", False, None
        both = status in ("Modified", "Unchanged")
        changed = bool(gd["changed"])
        if name == "_geom_status":
            if not both:
                return "—", False, None
            return ("Changed" if changed else "Same"), changed, None
        if name == f"_added_{unit}":
            v = gd["added_m"]
            return _fmt_num(v), bool(v) and both, None
        if name == f"_removed_{unit}":
            v = gd["removed_m"]
            return _fmt_num(v), bool(v) and both, None
        if name == "_changed_pct":
            if gd["pct"] is None:
                return "", False, None
            return f"{gd['pct']:.2f}%", changed, None
        if name == "_parts_changed":
            if not both:
                return "—", False, None
            n = gd.get("parts_changed", 0)
            return f"{n:,}", n > 0, None
        if name == "_parts":
            po, pn = gd["parts_old"], gd["parts_new"]
            if not both:
                return str(pn if status == "Added" else po), False, None
            if po != pn:
                return f"{po:,} → {pn:,}", True, (str(po), str(pn))
            return f"{pn:,}", False, None
        return "", False, None

    def _display_results(self, old_features, new_features,
                         display_fields, check_fields,
                         old_name, new_name,
                         geom_diffs=None, gtype=None, progress_base=45):
        """Populate the results table with colour-coded comparison data."""
        tbl = self.results_table
        t = self.theme
        geom_diffs = geom_diffs or {}
        extra_names = (self._geom_extra_names(gtype, self._geom_by_part_run)
                       if geom_diffs else [])
        unit = _geom_unit(gtype) if geom_diffs else ""
        all_ids = set(old_features.keys()) | set(new_features.keys())
        total = len(all_ids)

        tbl.setSortingEnabled(False)
        tbl.setRowCount(total)
        tbl.setColumnCount(
            len(display_fields) + len(extra_names) + COL_FIRST_FIELD)
        tbl.setHorizontalHeaderLabels(
            ["Change", "Decision"] + display_fields + extra_names
        )
        extra_tips = {
            "_geom_status": "Exact geometry comparison (within the "
                            "tolerance): Same or Changed",
            f"_added_{unit}": f"{unit.capitalize()} present in NEW but not "
                              f"in OLD",
            f"_removed_{unit}": f"{unit.capitalize()} present in OLD but "
                               f"not in NEW",
            "_changed_pct": f"(added + removed) ÷ OLD {unit} × 100",
            "_parts": "Number of parts in the geometry (OLD → NEW when it "
                      "differs)",
            "_parts_changed": "Parts that were added, removed or modified "
                              "(each has a review point)",
        }
        for c, name in enumerate(extra_names,
                                 COL_FIRST_FIELD + len(display_fields)):
            hi = tbl.horizontalHeaderItem(c)
            if hi:
                hi.setToolTip(extra_tips.get(name, ""))

        counts = {s: 0 for s in STATUSES}
        value_total = 0
        geom_change_count = 0
        geom_field_set = set(self._get_geom_field_names())
        check_set = set(check_fields)

        for row, fid in enumerate(sorted(all_ids, key=str)):
            in_old = fid in old_features
            in_new = fid in new_features
            changed_fields = set()
            gd = geom_diffs.get(fid)

            if in_old and in_new:
                old_d, new_d = old_features[fid], new_features[fid]
                changed_fields = {
                    fn for fn in display_fields
                    if fn in check_set
                    and not self.values_equal(old_d.get(fn), new_d.get(fn))
                }
                geom_changed = False
                if gd is not None:
                    # the tolerance-aware result decides, so summary numbers
                    # that differ only by noise are not flagged
                    geom_changed = bool(gd["changed"])
                    if not geom_changed:
                        changed_fields -= geom_field_set
                elif changed_fields & geom_field_set:
                    geom_changed = True
                if geom_changed:
                    geom_change_count += 1
                status = ("Modified" if changed_fields or geom_changed
                          else "Unchanged")
                data = new_d
            elif in_new:
                status = "Added"
                data = new_features[fid]
            else:
                status = "Deleted"
                data = old_features[fid]

            counts[status] += 1

            tbl.setVerticalHeaderItem(row, QTableWidgetItem(str(row + 1)))

            status_item = self._new_row_item(status, status)
            status_item.setData(_ROLE_USER, fid)
            status_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            tbl.setItem(row, COL_STATUS, status_item)

            dec_item = self._new_row_item("", status)
            if status in REVIEWABLE:
                self._style_decision_item(dec_item, None, status)
            else:
                dec_item.setText("—")
                dec_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            tbl.setItem(row, COL_DECISION, dec_item)

            for col, fn in enumerate(display_fields, COL_FIRST_FIELD):
                if fn in changed_fields:
                    old_v = self.format_value(old_d.get(fn))
                    new_v = self.format_value(new_d.get(fn))
                    item = _SortItem(f"{old_v} → {new_v}")
                    value_total += 1
                    item.setData(_ROLE_USER, (old_v, new_v))
                    item.setData(_ROLE_KIND, "changed")
                    item.setBackground(QBrush(t.bg("ChangedField")))
                    item.setForeground(QBrush(QColor(t.row_text)))
                    item.setToolTip(
                        f"{fn}\nOLD: {old_v or '(empty)'}\n"
                        f"NEW: {new_v or '(empty)'}\n"
                        "Right-click to accept or reject")
                else:
                    item = self._new_row_item(
                        self.format_value(data.get(fn, '')), status)
                tbl.setItem(row, col, item)

            for col, name in enumerate(
                    extra_names, COL_FIRST_FIELD + len(display_fields)):
                text, pink, pair = self._geom_extra_cell(
                    name, status, gd, unit)
                item = self._new_row_item(text, status)
                if pink:
                    item.setData(_ROLE_KIND, "changed")
                    item.setBackground(QBrush(t.bg("ChangedField")))
                if pair:
                    item.setData(_ROLE_USER, pair)
                    value_total += pink
                tbl.setItem(row, col, item)

            if total > 0 and row % 50 == 0:
                self.progress_bar.setValue(
                    progress_base
                    + int((95 - progress_base) * (row + 1) / total))
                QApplication.processEvents()

        tbl.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.ResizeToContents
        )
        tbl.setSortingEnabled(True)

        self._has_results = True
        self._reviewable_total = sum(counts[s] for s in REVIEWABLE)
        self._value_total = value_total
        self.results_stack.setCurrentIndex(1)
        self.export_btn.setEnabled(True)
        for w in self.review_widgets:
            w.setEnabled(True)
        self.map_btn.setEnabled(bool(geom_diffs))

        for status in STATUSES:
            self._set_filter_count(status, counts[status])
        self.geom_summary_label.setText(
            f"{geom_change_count:,} geometry change(s)"
            if (geom_field_set or geom_diffs) else "")

        # Field picker for the value filter (same icons/tooltips as Join field)
        previous = self.filter_field_combo.currentText()
        self.filter_field_combo.blockSignals(True)
        self.filter_field_combo.clear()
        self.filter_field_combo.addItem(self._all_fields_icon(), "All fields", -1)
        for col, fn in enumerate(display_fields + extra_names,
                                 COL_FIRST_FIELD):
            icon, tip = self._field_icon_and_tip(fn)
            self.filter_field_combo.addItem(icon, fn, col)
            self.filter_field_combo.setItemData(
                self.filter_field_combo.count() - 1, tip, _ROLE_TOOLTIP)
        idx = self.filter_field_combo.findText(previous)
        self.filter_field_combo.setCurrentIndex(max(idx, 0))
        self.filter_field_combo.blockSignals(False)
        self._refresh_value_suggestions()

        self._compare_info = {
            "old_name": old_name,
            "new_name": new_name,
            "join_field": self.join_field_combo.currentText(),
            "checked_fields": [f for f in check_fields
                               if f in display_fields],
            "geometry": sorted(geom_field_set),
            "tolerance": self.geom_tol_spin.value() if geom_diffs else None,
            "analysis": ("Singlepart (each part)" if self._geom_by_part_run
                         else "Multipart (whole feature)"
                         if self.geom_mode_combo.isEnabled()
                         else "Singlepart input")
            if geom_diffs else None,
            "review": self._review_note if geom_diffs else None,
            "review_path": self._review_saved if geom_diffs else "",
            "counts": dict(counts),
        }

        self.apply_filters()

        self._set_status(
            f"OLD ({old_name}) → NEW ({new_name}):  "
            f"{counts['Added']:,} added · {counts['Deleted']:,} deleted · "
            f"{counts['Modified']:,} modified · "
            f"{counts['Unchanged']:,} unchanged"
            + (f"  ·  {self._review_note}" if self._review_note else ""),
            "success"
        )
