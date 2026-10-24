# -*- coding: utf-8 -*-
"""
Export helpers shared by Attribute Explorer and Pivot Table.

Every export records its scope (layer, records, filters, time) so a file
can always be traced back: Excel workbooks get a "Scope" sheet and CSV
files a ``<name>_scope.txt`` next to them.  Values are written raw (full
precision, ISO dates, NULL as an empty cell).
"""

import csv
import datetime
import os
import re

from . import fields as F

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font as XlFont, PatternFill, Alignment
    HAS_OPENPYXL = True
except ImportError:  # pragma: no cover - depends on the QGIS install
    HAS_OPENPYXL = False

XLSX_MAX_ROWS = 1048575          # Excel sheet limit minus the header row


def export_value(value):
    """Plain, lossless text for CSV."""
    if F.is_null(value):
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return repr(value)
    if isinstance(value, (int, str)):
        return str(value)
    return F.display_value(value)


def excel_value(value):
    """Typed cell value for openpyxl (numbers stay numbers)."""
    if F.is_null(value):
        return None
    if isinstance(value, (bool, int, float, str)):
        if isinstance(value, float) and value != value:     # NaN
            return None
        return value
    group = F.group_value(value)
    if isinstance(group, (datetime.date, datetime.datetime, datetime.time)):
        if isinstance(group, datetime.datetime) and group.tzinfo is not None:
            return group.replace(tzinfo=None)
        return group
    return F.display_value(value)


def safe_field_names(names, max_len=63):
    """Unique, GeoPackage-safe field names (letters, digits, underscore)."""
    out, seen = [], set()
    for name in names:
        base = re.sub(r"[^0-9A-Za-z_]+", "_", str(name)).strip("_") or "field"
        if base[0].isdigit():
            base = "f_" + base
        base = base[:max_len]
        candidate, n = base, 2
        while candidate.lower() in seen:
            suffix = f"_{n}"
            candidate = base[:max_len - len(suffix)] + suffix
            n += 1
        seen.add(candidate.lower())
        out.append(candidate)
    return out


def scope_lines(meta):
    """'key: value' lines for a scope record (dict, ordered)."""
    stamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    lines = [f"Exported: {stamp}", "Tool: GeoTable Studio"]
    lines += [f"{k}: {v}" for k, v in meta.items()]
    return lines


def write_csv(path, header, rows, meta=None):
    """Write *rows* (iterables of raw values); returns the row count and
    writes the scope next to the file when *meta* is given."""
    count = 0
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(header)
        for row in rows:
            writer.writerow([export_value(v) for v in row])
            count += 1
    if meta is not None:
        base, _ext = os.path.splitext(path)
        with open(base + "_scope.txt", "w", encoding="utf-8") as fh:
            fh.write("\n".join(scope_lines(meta)) + "\n")
    return count


def write_xlsx(path, sheet_title, header, rows, meta=None, header_groups=None):
    """Write one data sheet (+ a Scope sheet); returns the row count."""
    if not HAS_OPENPYXL:
        raise RuntimeError("Excel export needs the 'openpyxl' Python package")
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(sheet_title[:31] or "Data")
    bold = XlFont(bold=True)
    fill = PatternFill(start_color="E4E4E4", end_color="E4E4E4",
                       fill_type="solid")

    def header_row(values):
        cells = []
        for v in values:
            cell = _cell(ws, v)
            cell.font = bold
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cells.append(cell)
        return cells

    widths = [max(8, min(50, len(str(h)) + 4)) for h in header]
    for i, w in enumerate(widths):
        ws.column_dimensions[_col_letter(i + 1)].width = w
    ws.freeze_panes = "A3" if header_groups else "A2"
    if header_groups:
        ws.append(header_row(header_groups))
    ws.append(header_row(header))
    count = 0
    for row in rows:
        if count >= XLSX_MAX_ROWS:
            raise RuntimeError(
                f"Excel holds at most {XLSX_MAX_ROWS:,} rows; use CSV instead")
        ws.append([excel_value(v) for v in row])
        count += 1
    if meta is not None:
        info = wb.create_sheet("Scope")
        info.column_dimensions["A"].width = 24
        info.column_dimensions["B"].width = 90
        for line in scope_lines(meta):
            key, _sep, value = line.partition(": ")
            k = _cell(info, key)
            k.font = bold
            info.append([k, value])
    wb.save(path)
    return count


def _cell(ws, value):
    from openpyxl.cell import WriteOnlyCell
    return WriteOnlyCell(ws, value=value)


def _col_letter(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)
