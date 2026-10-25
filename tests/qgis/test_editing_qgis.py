# -*- coding: utf-8 -*-
"""Attribute Explorer editing, forms, bulk tools and browse parity, in real
QGIS (run with the QGIS Python: python tests/qgis/test_editing_qgis.py)."""

import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import harness as H  # noqa: E402

APP = H.start_qgis()
H.plugin_package()

from qgis.PyQt import sip  # noqa: E402
from qgis.PyQt.QtCore import QEvent, QPoint, Qt  # noqa: E402
from qgis.PyQt.QtGui import QColor, QKeyEvent, QKeySequence  # noqa: E402
from qgis.PyQt.QtWidgets import QMenu, QMessageBox  # noqa: E402
from qgis.core import (Qgis, QgsConditionalStyle, QgsFeatureRequest, QgsField,  # noqa: E402
                       QgsFieldConstraints, QgsRectangle)

from attribute_table_studio import classFactory  # noqa: E402
from attribute_table_studio.core import fields as F  # noqa: E402
from attribute_table_studio.core import filters as FL  # noqa: E402
from attribute_table_studio.core import settings  # noqa: E402
from attribute_table_studio.core import subset as SUB  # noqa: E402
from attribute_table_studio.ui import explorer_dialogs as D  # noqa: E402
from attribute_table_studio.ui import explorer_edit as E  # noqa: E402
from attribute_table_studio.ui import icons  # noqa: E402

for _name in ("information", "warning", "critical"):
    setattr(QMessageBox, _name, staticmethod(lambda *a, **k: None))


class Asker:
    """Scripted answers for the Explorer prompts (records what was asked)."""

    def __init__(self):
        self.answers, self.asked = [], []

    def __call__(self, parent, title, text, choices, default=None, icon="question"):
        self.asked.append((title, text, [c[0] for c in choices]))
        if self.answers:
            return self.answers.pop(0)
        return default or choices[-1][0]


ASK = Asker()
E.ASK = ASK


def idle(panel):
    return panel._has_result and not panel.runner.busy and not panel._data_timer.isActive() \
        and not panel._mutation_timer.isActive()


def settle(panel):
    assert H.wait_for(lambda: idle(panel), 60)
    H.settle(0.05)
    assert H.wait_for(lambda: idle(panel), 60)


class EditingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.iface = H.FakeIface()
        cls.plugin = classFactory(cls.iface)
        cls.plugin.initGui()

    @classmethod
    def tearDownClass(cls):
        cls.plugin.unload()

    def setUp(self):
        ASK.answers, ASK.asked = [], []
        self.layer = H.demo_layer(n=400, seed=21, name=f"Edit parcels {self.id()[-12:]}")
        self.iface.setActiveLayer(self.layer)
        self.panel = self.plugin.open_explorer(self.layer)
        settle(self.panel)
        self.fl = self.layer.fields().indexOf("SCORE_MAX")
        self.lu = self.layer.fields().indexOf("LAND_USE")

    def tearDown(self):
        layer = self.layer
        if layer.isEditable():
            layer.rollBack()
        host = self.plugin.explorers.get(layer.id())
        if host is not None:
            host.close()
        H.settle(0.1)

    def start(self):
        self.panel.edit_bar.edit_btn.click()
        self.assertTrue(self.layer.isEditable())
        H.settle(0.15)

    # ------------------------------------------------------------------
    def test_01_rows_icons_and_quick_update_visibility(self):
        p, b = self.panel, self.panel.edit_bar
        top = p.columns_btn.mapTo(p, QPoint(0, 0)).y()
        self.assertGreater(b.mapTo(p, QPoint(0, 0)).y(), top)
        # View and Select sit before Columns, layer info stays left
        x = lambda w: w.mapTo(p, QPoint(0, 0)).x()
        self.assertLess(x(p.view_btn), x(p.select_btn))
        self.assertLess(x(p.select_btn), x(p.columns_btn))
        self.assertLess(x(p.layer_combo), x(p.view_btn))
        # both rows end at the same right edge (Dock / More)
        right = lambda w: w.mapTo(p, QPoint(w.width(), 0)).x()
        self.assertLessEqual(abs(right(p.dock_btn) - right(b.more_btn)), 2)
        self.assertNotEqual(icons.CALCULATE, icons.QUICK_UPDATE)
        self.assertEqual(icons.QUICK_UPDATE, "/mActionProcessSelected.svg")
        # the scope line sits on the left of the edit row (no row of its own);
        # the edit status is in the status bar, dim while read only
        y = lambda w: w.mapTo(p, QPoint(0, w.height() // 2)).y()
        self.assertLessEqual(abs(y(p.scope_label) - y(b.edit_btn)), 2)
        self.assertLess(x(p.scope_label), x(b.edit_btn))
        self.assertTrue(p.scope_label.fullText().startswith("Statistics for:"))
        self.assertLess(p.edit_bar.mapTo(p, QPoint(0, b.height())).y(),
                        p.view.mapTo(p, QPoint(0, 0)).y())
        extra = p.status.extra
        self.assertEqual(extra.fullText(), "Read only")
        self.assertEqual(extra.property("level"), "info")                # dim
        for btn in (b.save_btn, b.discard_btn, b.undo_btn, b.add_btn,
                    b.quick_btn, b.multi_btn, b.new_field_btn):
            self.assertFalse(btn.isEnabled(), btn.toolTip())
            self.assertIn("Unavailable", btn.toolTip())
        # QGIS's Field Calculator button is never disabled (it starts editing)
        self.assertTrue(b.calc_btn.isEnabled())
        self.assertIn("QGIS", b.calc_btn.toolTip())
        self.assertTrue(b.edit_btn.isEnabled())
        self.assertTrue(b.reload_btn.isEnabled())
        self.assertFalse(p.quick_bar.isVisibleTo(p))
        self.start()
        self.assertTrue(extra.fullText().startswith("Editing layer  •  No unsaved changes"))
        self.assertEqual(extra.property("level"), "editing-clean")      # orange
        self.assertTrue(b.quick_btn.isEnabled())
        self.assertFalse(b.quick_btn.isChecked())
        self.assertFalse(p.quick_bar.isVisibleTo(p))     # editing alone: hidden
        b.quick_btn.click()
        self.assertTrue(p.quick_bar.isVisibleTo(p))
        b.quick_btn.click()
        self.assertFalse(p.quick_bar.isVisibleTo(p))
        b.quick_btn.click()
        b.edit_btn.click()                                # stop (no changes)
        self.assertFalse(self.layer.isEditable())
        self.assertFalse(b.quick_btn.isChecked())
        self.assertFalse(p.quick_bar.isVisibleTo(p))

    def test_02_shared_session_and_stop_prompt(self):
        p, b = self.panel, self.panel.edit_bar
        self.layer.startEditing()                         # elsewhere in QGIS
        H.settle(0.2)
        self.assertTrue(b.edit_btn.isChecked())
        self.assertIn("Editing layer", p.status.extra.fullText())
        fid = p.model.fid(0)
        self.layer.changeAttributeValue(fid, self.lu, "OTHER")   # e.g. the map
        settle(p)
        self.assertIn("1 unsaved cell", p.status.extra.fullText())
        self.assertEqual(p.status.extra.property("level"), "editing")
        ASK.answers = ["cancel"]
        b.edit_btn.click()
        self.assertTrue(self.layer.isEditable())
        self.assertEqual(ASK.asked[-1][2], ["save", "discard", "cancel"])
        ASK.answers = ["discard"]
        b.edit_btn.click()
        self.assertFalse(self.layer.isEditable())
        self.assertNotEqual(self.layer.getFeature(fid)["LAND_USE"], "OTHER")

    def test_03_cell_edit_undo_redo_save(self):
        p = self.panel
        self.start()
        row = 0
        fid = p.model.fid(row)
        idx = p.model.index(row, self.fl)
        self.assertTrue(p.model.flags(idx) & Qt.ItemFlag.ItemIsEditable)
        self.assertTrue(p.model.setData(idx, 12.5))
        self.assertEqual(self.layer.getFeature(fid)["SCORE_MAX"], 12.5)
        settle(p)
        self.assertTrue(p.model.is_dirty(fid, self.fl))
        self.assertTrue(p.edit_bar.save_btn.isEnabled())
        self.assertEqual(p.edit.counts()["cells"], 1)
        p.edit.undo()
        self.assertNotEqual(self.layer.getFeature(fid)["SCORE_MAX"], 12.5)
        p.edit.redo()
        self.assertEqual(self.layer.getFeature(fid)["SCORE_MAX"], 12.5)
        self.assertTrue(p.save_edits())
        self.assertTrue(self.layer.isEditable())             # stays editing
        self.assertFalse(self.layer.isModified())
        saved = next(self.layer.dataProvider().getFeatures(
            QgsFeatureRequest().setFilterFid(fid)))
        self.assertEqual(saved["SCORE_MAX"], 12.5)
        settle(p)
        self.assertFalse(p.model.is_dirty(fid, self.fl))
        # profiles use the saved / unsaved values (12.5 is the new minimum)
        self.assertTrue(H.wait_for(lambda: self.fl in p.profiles, 30))
        self.assertEqual(p.profiles[self.fl].minimum, 12.5)

    def test_04_constraints_and_read_only_fields(self):
        p = self.panel
        layer = self.layer
        layer.setFieldConstraint(self.lu, QgsFieldConstraints.Constraint.ConstraintNotNull,
                                 QgsFieldConstraints.ConstraintStrength.ConstraintStrengthHard)
        layer.addExpressionField('"SCORE_MAX" * 2', QgsField("DOUBLE_MAX", D._MT.Double))
        settle(p)
        self.start()
        rejected = []
        p.model.editRejected.connect(rejected.append)
        idx = p.model.index(0, self.lu)
        before = p.model.raw_value(0, self.lu)
        self.assertFalse(p.model.setData(idx, None))
        self.assertTrue(rejected and "LAND_USE" in rejected[-1])
        self.assertEqual(p.model.raw_value(0, self.lu), before)
        virtual = layer.fields().indexFromName("DOUBLE_MAX")
        fid_idx = layer.fields().indexFromName("fid")
        self.assertFalse(p.edit.field_editable(virtual)[0])
        self.assertIn("Virtual", p.edit.field_editable(virtual)[1])
        self.assertFalse(p.edit.field_editable(fid_idx)[0])
        self.assertNotIn(virtual, p.model.editable_cols)
        self.assertNotIn(fid_idx, p.model.editable_cols)
        layer.removeExpressionField(virtual)

    def test_05_save_failure_keeps_edits(self):
        p = self.panel
        self.start()
        p.model.setData(p.model.index(0, self.fl), 1.25)
        layer = self.layer
        layer.commitChanges = lambda *a: False
        layer.commitErrors = lambda: ["ERROR: simulated provider failure"]
        try:
            self.assertFalse(p.save_edits())
        finally:
            del layer.commitChanges
            del layer.commitErrors
        self.assertTrue(layer.isEditable())
        self.assertTrue(layer.isModified())
        self.assertIn("Save failed", p.status.extra.fullText())
        self.assertEqual(p.status.extra.property("level"), "error")
        self.assertIn("simulated provider failure", ASK.asked[-1][1])

    def test_06_add_delete_and_edited_scope(self):
        p = self.panel
        self.start()
        n = self.layer.featureCount()
        p.add_record_in_table()
        self.assertEqual(self.layer.featureCount(), n + 1)
        self.assertTrue(p.status.action_btn.isVisibleTo(p))
        self.assertEqual(p.status.action_btn.text(), "Show edited and new")
        p.status.action_btn.click()
        settle(p)
        self.assertEqual(p.view_scope, "edited")
        self.assertEqual(len(p.model.fids), 1)
        self.assertLess(p.model.fids[0], 0)                  # temporary id
        p.set_view_scope("all")
        settle(p)
        # delete selected, some of them hidden by a filter
        p.header.cells[self.lu].set_filter(FL.ColumnFilter("LAND_USE", F.TEXT, "eq",
                                                           "RESIDENTIAL"))
        p.refresh()
        settle(p)
        shown = p.model.fids[:2]
        hidden = [f for f in self.layer.allFeatureIds() if f not in set(p.model.fids)][:1]
        self.layer.selectByIds(shown + hidden)
        ASK.answers = ["delete"]
        p.delete_selected()
        self.assertIn("1 of them are not shown", ASK.asked[-1][1])
        self.assertEqual(self.layer.featureCount(), n + 1 - 3)
        p.edit.undo()                                        # one undo step
        self.assertEqual(self.layer.featureCount(), n + 1)

    def test_07_feature_clipboard(self):
        p = self.panel
        self.layer.selectByIds(p.model.fids[:2])
        H.settle(0.1)
        self.assertTrue(p.edit_bar.copy_btn.isEnabled())      # browsing too
        self.assertFalse(p.edit_bar.paste_btn.isEnabled())
        p.copy_features()
        self.start()
        n = self.layer.featureCount()
        p.paste_features()
        self.assertEqual(self.layer.featureCount(), n + 2)
        self.layer.selectByIds(p.model.fids[:1])
        H.settle(0.1)
        p.cut_features()
        self.assertEqual(self.layer.featureCount(), n + 1)

    def test_08_new_and_delete_fields(self):
        p = self.panel
        cells = p.header.cells
        cells[self.lu].set_filter(FL.ColumnFilter("LAND_USE", F.TEXT, "eq", "RESIDENTIAL"))
        p.refresh()
        settle(p)
        self.start()
        dlg = D.NewFieldDialog(p)
        self.assertTrue(dlg.physical.isEnabled())
        dlg.name_edit.setText("NOTES")
        dlg.type_combo.setCurrentIndex(max(0, dlg.type_combo.findText("Text (string)")))
        dlg._apply()
        settle(p)
        self.assertGreaterEqual(self.layer.fields().indexFromName("NOTES"), 0)
        # unrelated filters survive the schema change
        lu = self.layer.fields().indexFromName("LAND_USE")
        self.assertIsNotNone(p.header.cells[lu].current_filter())
        dlg = D.NewFieldDialog(p)
        dlg.virtual.setChecked(True)
        dlg.name_edit.setText("HALF_MAX")
        dlg.expr.setExpression('"SCORE_MAX" / 2')
        dlg._apply()
        settle(p)
        virtual = self.layer.fields().indexFromName("HALF_MAX")
        self.assertEqual(self.layer.fields().fieldOrigin(virtual),
                         Qgis.FieldOrigin.Expression)
        ASK.answers = ["delete"]
        dlg = D.DeleteFieldDialog(p, "NOTES")
        dlg._apply()
        settle(p)
        self.assertLess(self.layer.fields().indexFromName("NOTES"), 0)
        ASK.answers = ["delete"]
        dlg = D.DeleteFieldDialog(p, "HALF_MAX")
        dlg._apply()
        self.assertLess(self.layer.fields().indexFromName("HALF_MAX"), 0)

    def test_09_field_calculator_is_qgis_dialog(self):
        """The button opens QGIS's own Field Calculator on this layer; it
        starts editing itself, pre-ticks "Only update selected" and uses
        @row_number - all QGIS's code."""
        from qgis.gui import QgsFieldCalculator
        from qgis.PyQt.QtWidgets import QCheckBox, QComboBox, QGroupBox
        p, layer = self.panel, self.layer
        self.assertFalse(layer.isEditable())
        ids = sorted(p.model.fids[:3])
        layer.selectByIds(ids)
        seen = {}

        def run(dlg):
            seen["title"] = dlg.windowTitle()
            seen["update"] = dlg.findChild(QGroupBox, "mUpdateExistingGroupBox").isChecked()
            seen["field"] = dlg.findChild(QComboBox, "mExistingFieldComboBox").currentText()
            seen["only"] = dlg.findChild(QCheckBox, "mOnlyUpdateSelectedCheckBox").isChecked()
            self._calc_builder(dlg).setExpressionText("@row_number * 10")
            dlg.accept()
            return 1

        original = QgsFieldCalculator.exec
        QgsFieldCalculator.exec = run
        try:
            p.open_calculator("ZONE_CODE")        # header menu: "Calculate ZONE_CODE…"
        finally:
            QgsFieldCalculator.exec = original
        self.assertIn(layer.name(), seen["title"])
        self.assertEqual((seen["update"], seen["field"], seen["only"]), (True, "ZONE_CODE", True))
        self.assertTrue(layer.isEditable())                     # QGIS started editing
        self.assertEqual([layer.getFeature(f)["ZONE_CODE"] for f in ids], [10, 20, 30])
        settle(p)
        self.assertTrue(p.edit_bar.edit_btn.isChecked())        # the Explorer follows
        self.assertTrue(p.can_calculate("ZONE_CODE"))
        layer.addExpressionField('"ZONE_CODE" * 2', QgsField("TWICE", D._MT.Int))
        self.assertFalse(p.can_calculate("TWICE"))             # QGIS never lists virtual fields
        layer.removeExpressionField(layer.fields().indexFromName("TWICE"))

    @staticmethod
    def _calc_builder(dlg):
        """The expression builder inside QGIS's Field Calculator (PyQGIS
        returns it as a plain QWidget)."""
        from qgis.gui import QgsExpressionBuilderWidget
        from qgis.PyQt.QtWidgets import QWidget
        widget = dlg.findChild(QWidget, "builder")
        assert widget.metaObject().className() == "QgsExpressionBuilderWidget"
        return sip.cast(widget, QgsExpressionBuilderWidget)

    def _quick(self, field, expression):
        p = self.panel
        if not p.edit_bar.quick_btn.isChecked():
            p.edit_bar.quick_btn.click()
        q = p.quick_bar
        q.field_combo.setField(field)
        q.expression.setExpression(expression)
        p._update_quick_state()
        return q

    def test_10_quick_update_all_and_selected(self):
        """Not filtered: "Update All" and "Update Selected", as in QGIS."""
        p, layer = self.panel, self.layer
        self.start()
        q = self._quick("SCORE_MAX", '"SCORE_MAX" + 1000')
        self.assertEqual(q.run_btn.text(), "Update All")
        self.assertFalse(q.selected_btn.isEnabled())            # nothing selected
        self.assertTrue(q.run_btn.isEnabled())
        ids = sorted(p.model.fids[:4])
        layer.selectByIds(ids)
        H.settle(0.1)
        self.assertTrue(q.selected_btn.isEnabled())
        before = {f: f_["SCORE_MAX"] for f_ in layer.getFeatures() for f in [f_.id()]}
        q.preview_btn.click()                                   # a toggle
        self.assertTrue(q.preview.isVisibleTo(p))
        q.preview_btn.click()
        self.assertFalse(q.preview.isVisibleTo(p))
        q.preview_btn.click()
        self.assertTrue(q.preview.isVisibleTo(p))
        self.assertIn(f"Update All: {layer.featureCount():,} records", q.preview_title.text())
        self.assertIn("Update Selected: 4", q.preview_title.text())
        self.assertEqual({f.id(): f["SCORE_MAX"] for f in layer.getFeatures()}, before)
        depth = layer.undoStack().count()
        q.selected_btn.click()
        self.assertEqual(layer.undoStack().count(), depth + 1)  # one undo step
        after = {f.id(): f["SCORE_MAX"] for f in layer.getFeatures()}
        for f, v in before.items():
            self.assertAlmostEqual(after[f], v + 1000 if f in ids else v)
        p.edit.undo()
        self.assertEqual({f.id(): f["SCORE_MAX"] for f in layer.getFeatures()}, before)
        # Update All: every feature, @row_number in the layer's order
        q = self._quick("ZONE_CODE", "@row_number")
        q.run_btn.click()
        values = [f["ZONE_CODE"] for f in layer.getFeatures()]
        self.assertEqual(values, list(range(1, layer.featureCount() + 1)))

    def test_11_multi_edit(self):
        p = self.panel
        self.start()
        ids = p.model.fids[:3]
        self.layer.selectByIds(ids)
        dlg = D.MultiEditDialog(p, ids)
        dlg.form.changeAttribute("LAND_USE", "MIXED")
        dlg._apply()
        for fid in ids:
            self.assertEqual(self.layer.getFeature(fid)["LAND_USE"], "MIXED")

    def test_12_form_view(self):
        p = self.panel
        fid = p.model.fid(5)
        p.view.setCurrentIndex(p.model.index(5, self.lu))
        p.set_form_view(True)
        self.assertTrue(p.is_form_view())
        self.assertEqual(p.form_view.current, fid)
        self.start()
        p.form_view.form.changeAttribute("LAND_USE", "FORM EDIT")
        p.set_form_view(False)                               # writes the form
        self.assertFalse(p.is_form_view())
        self.assertEqual(self.layer.getFeature(fid)["LAND_USE"], "FORM EDIT")
        settle(p)
        row = p.model.row_of(fid)
        self.assertEqual(p.view.currentIndex().row(), row)

    def test_13_scopes(self):
        p = self.panel
        layer = self.layer
        # selected scope and the Selected only button are one setting
        layer.selectByIds(p.model.fids[:5])
        p.selected_btn.click()
        settle(p)
        self.assertEqual(p.view_scope, "selected")
        self.assertEqual(len(p.model.fids), 5)
        p.set_view_scope("all")
        self.assertFalse(p.selected_btn.isChecked())
        settle(p)
        # visible on map
        canvas = self.iface.mapCanvas()
        canvas.setDestinationCrs(layer.crs())
        ext = layer.extent()
        canvas.setExtent(QgsRectangle(ext.xMinimum(), ext.yMinimum(),
                                      ext.center().x(), ext.center().y()))
        p.set_view_scope("map")
        settle(p)
        self.assertGreater(len(p.model.fids), 0)
        self.assertLess(len(p.model.fids), layer.featureCount())
        self.assertIn("visible on the map", p.scope_label.fullText())
        # failing constraints
        layer.setFieldConstraint(self.lu, QgsFieldConstraints.Constraint.ConstraintExpression,
                                 QgsFieldConstraints.ConstraintStrength.ConstraintStrengthSoft)
        layer.setConstraintExpression(self.lu, '"LAND_USE" <> \'RURAL\'')
        p.set_view_scope("failing")
        settle(p)
        rural = sum(1 for f in layer.getFeatures() if f["LAND_USE"] == "RURAL")
        self.assertEqual(len(p.model.fids), rural)
        self.assertIn(f"(0 hard, {rural} soft)", p.scope_label.fullText())

    def test_14_expression_filter_and_stored(self):
        p = self.panel
        total = self.layer.featureCount()
        p.set_expression_filter('"SCORE_MAX" > 80')
        settle(p)
        a = len(p.model.fids)
        self.assertLess(a, total)
        self.assertIn("Expression", p.scope_label.toolTip())
        self.assertTrue(p.expr_chip.isVisibleTo(p))
        p.set_expression_filter('"LAND_USE" = \'RESIDENTIAL\'', "and")
        settle(p)
        self.assertLessEqual(len(p.model.fids), a)
        p.set_expression_filter('"SCORE_MAX" < 10', "or")
        settle(p)
        manager = self.layer.storedExpressionManager()
        dlg = D.StoredExpressionsDialog(p)
        before = len(manager.storedExpressions())
        manager.addStoredExpression("high", '"SCORE_MAX" > 90')
        dlg._fill()
        self.assertEqual(dlg.list.topLevelItemCount(), before + 1)
        dlg.list.setCurrentItem(dlg.list.topLevelItem(dlg.list.topLevelItemCount() - 1))
        dlg._apply()
        p.set_expression_filter(dlg.applied)
        settle(p)
        self.assertEqual(p.expr_filter, '"SCORE_MAX" > 90')
        p.clear_filters()
        settle(p)
        self.assertIsNone(p.expr_filter)
        self.assertEqual(len(p.model.fids), total)

    def test_15_select_menu_sort_and_selected_top(self):
        p = self.panel
        layer = self.layer
        p.select_layer_all()
        self.assertEqual(layer.selectedFeatureCount(), layer.featureCount())
        p.deselect_all()
        layer.selectByIds(p.model.fids[:2])
        p.invert_layer_selection()
        self.assertEqual(layer.selectedFeatureCount(), layer.featureCount() - 2)
        p.deselect_all()

        def menu_texts():
            p.select_menu.aboutToShow.emit()
            return [a.text() for a in p.select_menu.actions() if not a.isSeparator()]
        shown = menu_texts()
        for text in ("Select All Features", "Invert Selection",
                     "Deselect Features from All Layers", "Select Features by Expression…",
                     "Select Features by Value…", "Pan to Selection", "Zoom to Selection"):
            self.assertIn(text, shown)
        # the whole-layer variants only differ - and only show - while the
        # table is filtered
        self.assertNotIn("Invert Selection (Whole Layer)", shown)
        self.assertNotIn("Select All Features (Whole Layer)", shown)
        p.header.cells[self.lu].set_filter(
            FL.ColumnFilter("LAND_USE", F.TEXT, "eq", "RESIDENTIAL"))
        p.refresh()
        settle(p)
        shown = menu_texts()
        self.assertIn("Invert Selection (Whole Layer)", shown)
        self.assertIn("Select All Features (Whole Layer)", shown)
        filtered = sorted(p.model.fids)
        layer.selectByIds(filtered[:2])
        H.settle(0.1)
        p.invert_selection()
        self.assertEqual(set(layer.selectedFeatureIds()), set(filtered[2:]))
        layer.selectByIds(filtered[:2])
        H.settle(0.1)
        p.invert_layer_selection()
        self.assertEqual(layer.selectedFeatureCount(), layer.featureCount() - 2)
        p.deselect_all()
        p.clear_filters()
        settle(p)
        # sort by expression (descending)
        p.sort_expr = ('"SCORE_MAX" * -1', False)
        p.refresh(sort_only=True)
        settle(p)
        values = [p.model.raw_value(r, self.fl) for r in range(5)]
        self.assertEqual(values, sorted(values, reverse=True))
        # selected rows to the top, filter membership unchanged
        n = len(p.model.fids)
        last = p.model.fids[-1]
        layer.selectByIds([last])
        p.set_selected_top(True)
        self.assertEqual(p.model.fids[0], last)
        self.assertEqual(len(p.model.fids), n)

    def test_16_conditional_formatting(self):
        p = self.panel
        style = QgsConditionalStyle('"SCORE_MAX" > 50')
        style.setBackgroundColor(QColor("#123456"))
        self.layer.conditionalStyles().setRowStyles([style])
        H.settle(0.1)
        rows = [r for r in range(min(40, len(p.model.fids)))]
        hits = [r for r in rows if (p.model.raw_value(r, self.fl) or 0) > 50]
        misses = [r for r in rows if (p.model.raw_value(r, self.fl) or 0) <= 50]
        bg = p.model.data(p.model.index(hits[0], self.lu), Qt.ItemDataRole.BackgroundRole)
        self.assertEqual(QColor(bg).name(), "#123456")
        self.assertIsNone(p.model.data(p.model.index(misses[0], self.lu),
                                       Qt.ItemDataRole.BackgroundRole))
        self.layer.conditionalStyles().setRowStyles([])

    def test_17_keyboard_ownership(self):
        p = self.panel
        self.start()
        view = p.view
        # a click focuses the cell without changing the feature selection
        self.layer.removeSelection()
        rect = view.visualRect(p.model.index(2, self.fl))
        from qgis.PyQt.QtTest import QTest
        QTest.mouseClick(view.viewport(), Qt.MouseButton.LeftButton, pos=rect.center())
        self.assertEqual(self.layer.selectedFeatureCount(), 0)
        self.assertEqual(view.currentIndex().row(), 2)
        p.model.setData(p.model.index(2, self.fl), 3.5)
        fid = p.model.fid(2)
        # Ctrl+Z belongs to the Explorer outside an editor (not QGIS)
        override = QKeyEvent(QEvent.Type.ShortcutOverride, Qt.Key.Key_Z,
                             Qt.KeyboardModifier.ControlModifier)
        self.assertTrue(view.event(override))
        self.assertTrue(override.isAccepted())
        QTest.keyClick(view, Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
        self.assertNotEqual(self.layer.getFeature(fid)["SCORE_MAX"], 3.5)
        # Ctrl+S saves the layer while editing (accepted at the panel)
        p.model.setData(p.model.index(2, self.fl), 4.5)
        save = QKeyEvent(QEvent.Type.ShortcutOverride, Qt.Key.Key_S,
                         Qt.KeyboardModifier.ControlModifier)
        self.assertTrue(p.event(save))
        QTest.keyClick(p, Qt.Key.Key_S, Qt.KeyboardModifier.ControlModifier)
        self.assertFalse(self.layer.isModified())

    def test_18_close_prompt_keeps_buffer(self):
        p = self.panel
        self.start()
        p.model.setData(p.model.index(0, self.fl), 2.5)
        host = self.plugin.explorers[self.layer.id()]
        ASK.answers = ["cancel"]
        host._window.close()
        H.settle(0.2)
        self.assertIn(self.layer.id(), self.plugin.explorers)
        ASK.answers = ["keep"]
        host._window.close()
        H.settle(0.3)
        self.assertNotIn(self.layer.id(), self.plugin.explorers)
        self.assertTrue(self.layer.isEditable())
        self.assertTrue(self.layer.isModified())             # nothing discarded

    def test_19_map_filter_frozen_while_editing(self):
        p = self.panel
        cell = p.header.cells[self.lu]
        cell.set_filter(FL.ColumnFilter("LAND_USE", F.TEXT, "eq", "RESIDENTIAL"))
        p.refresh()
        settle(p)
        p.query_btn.setChecked(True)
        settle(p)
        frozen = self.layer.subsetString()
        self.assertIn("RESIDENTIAL", frozen)
        self.start()
        self.assertFalse(p.query_btn.isEnabled())
        self.assertIn("frozen", p.query_btn.toolTip())
        cell.set_filter(FL.ColumnFilter("LAND_USE", F.TEXT, "eq", "INDUSTRIAL"))
        p.refresh()
        settle(p)
        self.assertEqual(self.layer.subsetString(), frozen)    # unchanged
        self.assertEqual(len(p.model.fids), 0)                 # only map records
        p.edit_bar.edit_btn.click()                            # stop (no changes)
        settle(p)
        self.assertIn("INDUSTRIAL", self.layer.subsetString())
        self.assertTrue(p.query_btn.isEnabled())
        p.query_btn.setChecked(False)

    def test_19b_map_filter_closed_while_editing(self):
        """QGIS refuses filter changes while editing: an explorer closed
        then gives the layer its own filter back once editing stops, and one
        reopened before that takes the map filter over."""
        p, layer = self.panel, self.layer
        p.header.cells[self.lu].set_filter(
            FL.ColumnFilter("LAND_USE", F.TEXT, "eq", "RESIDENTIAL"))
        p.refresh()
        settle(p)
        p.query_btn.setChecked(True)
        settle(p)
        frozen = layer.subsetString()
        self.assertIn("RESIDENTIAL", frozen)
        self.start()
        ASK.answers = ["keep"]
        p.model.setData(p.model.index(0, self.fl), 12.5)
        self.plugin.explorers[layer.id()].close()            # keep the layer edits
        H.settle(0.2)
        self.assertTrue(layer.isEditable())
        self.assertEqual(layer.subsetString(), frozen)       # cannot change yet
        # reopened while still editing: the button is on (frozen) with the
        # filters that made it, so the table matches the map
        p = self.plugin.open_explorer(layer)
        settle(p)
        self.assertTrue(p.query_btn.isChecked())
        self.assertFalse(p.query_btn.isEnabled())
        self.assertEqual(p.header.cells[self.lu].current_filter().value, "RESIDENTIAL")
        self.assertEqual(len(p.model.fids), layer.featureCount())
        ASK.answers = ["keep"]
        self.plugin.explorers[layer.id()].close()
        H.settle(0.2)
        self.assertEqual(layer.subsetString(), frozen)
        layer.rollBack()                                     # editing stops
        H.settle(0.2)
        self.assertEqual(layer.subsetString(), "")
        self.assertIsNone(SUB.held(layer))

    def test_20_more_and_header_menus(self):
        p = self.panel
        self.start()
        shown = []
        original = QMenu.exec
        QMenu.exec = lambda menu, *a: shown.extend(
            act.text() for act in menu.actions() if not act.isSeparator())
        try:
            p._fill_more_menu()
            texts = [a.text() for a in p.edit_bar.more_menu.actions()]
            p._header_menu(QPoint(p.header.sectionViewportPosition(self.fl) + 5, 40))
        finally:
            QMenu.exec = original
        for text in ("Open Feature Form…", "Layer Feature Actions", "Copy Selected Rows as Text",
                     "Copy Cell Value", "Set Cell to NULL", "Review Unsaved Changes…",
                     "Layer Field Properties…", "Open Native Attribute Table", "Help"):
            self.assertIn(text, texts)
        for text in ("Sort SCORE_MAX Ascending", "Autosize This Column",
                     "Set Column Width…", "Calculate SCORE_MAX…"):
            self.assertIn(text, shown)

    def test_22_cell_editor_widgets(self):
        from qgis.PyQt.QtTest import QTest
        from qgis.PyQt.QtWidgets import QAbstractItemView
        p, view, layer = self.panel, self.panel.view, self.layer
        layer.setFieldConstraint(self.lu, QgsFieldConstraints.Constraint.ConstraintNotNull,
                                 QgsFieldConstraints.ConstraintStrength.ConstraintStrengthHard)
        self.start()
        delegate = view.itemDelegate()
        idx = p.model.index(1, self.lu)
        fid = p.model.fid(1)
        before = layer.getFeature(fid)["LAND_USE"]
        view.setCurrentIndex(idx)
        view.edit(idx)
        H.settle(0.05)
        editor = view.indexWidget(idx)
        self.assertIsNotNone(editor)                 # the field's QGIS widget
        wrapper = delegate._wrappers[sip.unwrapinstance(editor)][0]
        self.assertEqual(wrapper.field().name(), "LAND_USE")
        wrapper.setValues(None, [])
        QTest.keyClick(editor, Qt.Key.Key_Return)    # hard constraint: stays open
        H.settle(0.05)
        self.assertEqual(view.state(), QAbstractItemView.State.EditingState)
        self.assertEqual(layer.getFeature(fid)["LAND_USE"], before)
        wrapper.setValues("TYPED", [])
        view.commit_editor()
        H.settle(0.1)
        self.assertNotEqual(view.state(), QAbstractItemView.State.EditingState)
        self.assertEqual(layer.getFeature(fid)["LAND_USE"], "TYPED")
        self.assertFalse(delegate._wrappers)         # wrapper released with the editor
        # Escape cancels only this cell
        view.edit(p.model.index(2, self.lu))
        H.settle(0.05)
        editor = view.indexWidget(p.model.index(2, self.lu))
        delegate._wrappers[sip.unwrapinstance(editor)][0].setValues("CANCELLED", [])
        QTest.keyClick(editor, Qt.Key.Key_Escape)
        H.settle(0.05)
        self.assertNotEqual(layer.getFeature(p.model.fid(2))["LAND_USE"], "CANCELLED")
        self.assertEqual(layer.getFeature(fid)["LAND_USE"], "TYPED")

    # ---- Quick update: QGIS's rules for what is updated --------------------
    def test_23_quick_update_filtered_without_query(self):
        """Table filtered, map filter (query) off: "Update Filtered" changes
        only the shown records, "Update Selected" every selected feature,
        and once the filters are cleared "Update All" reaches every record."""
        p, layer = self.panel, self.layer
        p.header.cells[self.lu].set_filter(FL.ColumnFilter("LAND_USE", F.TEXT, "eq",
                                                           "RESIDENTIAL"))
        p.refresh()
        settle(p)
        self.start()
        shown = set(p.model.fids)
        self.assertLess(len(shown), layer.featureCount())
        q = self._quick("ZONE_CODE", "99")
        self.assertEqual(q.run_btn.text(), "Update Filtered")
        q.run_btn.click()
        zone = {f.id(): f["ZONE_CODE"] for f in layer.getFeatures()}
        self.assertTrue(all(zone[f] == 99 for f in shown))
        self.assertTrue(all(zone[f] != 99 for f in zone if f not in shown))
        hidden = next(f for f in zone if f not in shown)
        layer.selectByIds([hidden])
        H.settle(0.1)
        q = self._quick("ZONE_CODE", "77")
        q.selected_btn.click()                       # selected but not shown: updated too
        self.assertEqual(layer.getFeature(hidden)["ZONE_CODE"], 77)
        p.clear_filters()
        settle(p)
        q = self._quick("ZONE_CODE", "55")
        self.assertEqual(q.run_btn.text(), "Update All")
        q.run_btn.click()
        self.assertEqual({f["ZONE_CODE"] for f in layer.getFeatures()}, {55})
        self.assertEqual(sum(1 for _ in layer.getFeatures()), layer.featureCount())

    def test_24_quick_update_with_query_on(self):
        """Map filter (query) on: only what the query shows can be edited -
        even "Update All" and QGIS's own calculator stop at the query."""
        from qgis.gui import QgsFieldCalculator
        from qgis.PyQt.QtWidgets import QCheckBox
        p, layer = self.panel, self.layer
        everything = set(layer.allFeatureIds())
        p.header.cells[self.lu].set_filter(FL.ColumnFilter("LAND_USE", F.TEXT, "eq",
                                                           "RESIDENTIAL"))
        p.refresh()
        settle(p)
        p.query_btn.setChecked(True)
        settle(p)
        queried = set(layer.allFeatureIds())                 # what the map shows
        self.assertLess(len(queried), len(everything))
        self.start()
        # a narrower table filter while editing: Update Filtered stays inside it
        fl = p.header.cells[self.fl]
        fl.set_filter(FL.ColumnFilter("SCORE_MAX", F.REAL, "gt", "80"))
        p.refresh()
        settle(p)
        narrower = set(p.model.fids)
        self.assertTrue(narrower < queried)
        q = self._quick("ZONE_CODE", "42")
        self.assertEqual(q.run_btn.text(), "Update Filtered")
        q.run_btn.click()
        changed = set(layer.editBuffer().changedAttributeValues().keys())
        self.assertEqual(changed, narrower)
        # no table filter left: Update All = every record the query shows
        p.clear_filters()
        settle(p)
        self.assertEqual(set(p.model.fids), queried)            # map filter frozen
        q = self._quick("ZONE_CODE", "43")
        self.assertEqual(q.run_btn.text(), "Update All")
        q.run_btn.click()
        changed = set(layer.editBuffer().changedAttributeValues().keys())
        self.assertEqual(changed, queried)                      # nothing outside the query
        # QGIS's own calculator (whole layer) also stops at the query
        layer.removeSelection()

        def run(dlg):
            dlg.findChild(QCheckBox, "mOnlyUpdateSelectedCheckBox").setChecked(False)
            self._calc_builder(dlg).setExpressionText("44")
            dlg.accept()
            return 1

        original = QgsFieldCalculator.exec
        QgsFieldCalculator.exec = run
        try:
            p.open_calculator("ZONE_CODE")
        finally:
            QgsFieldCalculator.exec = original
        changed = set(layer.editBuffer().changedAttributeValues().keys())
        self.assertEqual(changed, queried)
        self.assertEqual({f["ZONE_CODE"] for f in layer.getFeatures()}, {44})
        layer.rollBack()
        p.query_btn.setChecked(False)

    def test_25_quick_update_converts_values_like_qgis(self):
        """Values are converted with QgsField.convertCompatible as QGIS
        does: halves round away from zero, failures become NULL, long text
        is cut to the field length."""
        from attribute_table_studio.ui import explorer_ops as OPS
        p, layer = self.panel, self.layer
        self.start()
        # a decimal field with a precision (GeoPackage keeps none for REAL, and
        # QGIS then rounds nothing - the same here)
        layer.addAttribute(QgsField("P2", D._MT.Double, "double", 10, 2))
        settle(p)
        fid = sorted(p.model.fids)[0]
        layer.selectByIds([fid])
        H.settle(0.1)
        cases = [("P2", "0.125", 0.13),                    # double(10,2); Python round: 0.12
                 ("SCORE_MAX", "0.125", 0.125),         # no precision: not rounded
                 ("ZONE_CODE", "2.5", 3),                   # integer; Python round: 2
                 ("ZONE_CODE", "'abc'", None),              # cannot convert: NULL
                 ("PARCEL_ID", "'ABCDEFGHIJKLMNOPQRSTUVWXYZ'", "ABCDEFGHIJKLMNOPQRST")]
        for field, expression, expected in cases:
            q = self._quick(field, expression)
            q.selected_btn.click()
            got = layer.getFeature(fid)[field]
            self.assertEqual(None if F.is_null(got) else got, expected, (field, expression))
        self.assertIn("could not be stored", p.status.message.fullText())
        fld = layer.fields().field("ZONE_CODE")
        self.assertEqual(OPS.convert_like_qgis(fld, -2.5), (-3, None))

    def test_26_quick_update_measures_like_qgis(self):
        """$area / $perimeter on the project's ellipsoid in the project's
        units, as QGIS's calculation bar does."""
        from qgis.core import QgsDistanceArea, QgsProject
        from attribute_table_studio.ui import explorer_ops as OPS
        project = QgsProject.instance()
        saved = (project.crs(), project.ellipsoid(), project.areaUnits(),
                 project.distanceUnits())
        project.setCrs(self.layer.crs())     # QGIS has no ellipsoid without a project CRS
        project.setEllipsoid("EPSG:7030")
        project.setAreaUnits(Qgis.AreaUnit.Hectares)
        project.setDistanceUnits(Qgis.DistanceUnit.Kilometers)
        try:
            p, layer = self.panel, self.layer
            self.start()
            layer.addAttribute(QgsField("AREA_HA", D._MT.Double, "REAL"))
            settle(p)
            ids = sorted(p.model.fids[:2])
            layer.selectByIds(ids)
            H.settle(0.1)
            self._quick("AREA_HA", "$area").selected_btn.click()
            da = QgsDistanceArea()
            da.setSourceCrs(layer.crs(), project.transformContext())
            da.setEllipsoid("EPSG:7030")
            for fid in ids:
                geom = layer.getFeature(fid).geometry()
                got = layer.getFeature(fid)["AREA_HA"]
                expected = da.convertAreaMeasurement(da.measureArea(geom),
                                                     Qgis.AreaUnit.Hectares)
                self.assertAlmostEqual(got, expected, places=9)
                self.assertNotAlmostEqual(got, geom.area() / 10000, places=7)  # not planar
            idx = layer.fields().indexFromName("AREA_HA")
            for fid, _old, new in OPS.evaluate_preview(layer, idx, "$perimeter", ids):
                geom = layer.getFeature(fid).geometry()
                expected = da.convertLengthMeasurement(da.measurePerimeter(geom),
                                                       Qgis.DistanceUnit.Kilometers)
                self.assertAlmostEqual(new, expected, places=9)
        finally:
            project.setCrs(saved[0])
            project.setEllipsoid(saved[1])
            project.setAreaUnits(saved[2])
            project.setDistanceUnits(saved[3])

    # ---- review round 8 ------------------------------------------------
    def test_27_tick_boxes_select_like_rows(self):
        """Browsing and editing: a tick click selects only that row, Ctrl adds
        or removes, Shift selects a range."""
        from qgis.PyQt.QtTest import QTest
        p, layer = self.panel, self.layer
        header = p.view.row_header
        ctrl = Qt.KeyboardModifier.ControlModifier
        shift = Qt.KeyboardModifier.ShiftModifier
        none = Qt.KeyboardModifier.NoModifier

        def tick(row, mods=none):
            y = header.sectionViewportPosition(row) + header.sectionSize(row) // 2
            QTest.mouseClick(header.viewport(), Qt.MouseButton.LeftButton, mods, QPoint(8, y))
            H.settle(0.05)
            return set(layer.selectedFeatureIds())

        fid = p.model.fid
        for editing in (False, True):
            if editing:
                self.start()
            layer.removeSelection()
            self.assertEqual(tick(1), {fid(1)})
            self.assertEqual(tick(3), {fid(3)})                 # replaces, not adds
            self.assertEqual(tick(5, ctrl), {fid(3), fid(5)})
            self.assertEqual(tick(3, ctrl), {fid(5)})           # Ctrl removes again
            # Shift: from the last clicked row (3) to this one
            self.assertEqual(tick(8, shift), {fid(r) for r in range(3, 9)})

    def test_28_card_tooltips_bars_and_brief_stats(self):
        from qgis.PyQt.QtGui import QHelpEvent
        from qgis.PyQt.QtWidgets import QApplication, QToolTip
        p = self.panel
        cards = p.header.cards
        self.assertTrue(H.wait_for(lambda: cards[self.fl].state == "ready"
                                   and cards[self.lu].state == "ready", 30))
        H.settle(0.2)

        def tip(card, pos):
            QToolTip.hideText()
            ev = QHelpEvent(QEvent.Type.ToolTip, pos, card.mapToGlobal(pos))
            QApplication.sendEvent(card, ev)
            return QToolTip.text()

        card = p.header.cards[self.fl]
        card.grab()                          # paints it (lays out the bars)
        rect, _i = card._bars[0]
        # anywhere in the bar's column - also above a short bar - shows its count
        for pos in (QPoint(int(rect.center().x()), int(rect.top()) + 1),
                    QPoint(int(rect.center().x()), int(rect.bottom()) - 1)):
            text = tip(card, pos)
            self.assertIn("records", text)
            self.assertNotIn("Click for full statistics", text)
        self.assertIn("Click for full statistics", tip(card, QPoint(6, 6)))   # title
        text_card = p.header.cards[self.lu]
        text_card.grab()
        row_rect, row_tip = text_card._rows[0]
        self.assertIn("records", tip(text_card, row_rect.center().toPoint()))
        self.assertIn("RESIDENTIAL", row_tip)
        self.assertIn("Click for full statistics", tip(text_card, QPoint(6, 6)))

    @staticmethod
    def _visible_details():
        from attribute_table_studio.ui.explorer_header import ProfileDetails
        from qgis.PyQt.QtWidgets import QApplication
        return [w for w in QApplication.topLevelWidgets()
                if isinstance(w, ProfileDetails) and w.isVisible()]

    def _open_details(self, logical):
        p = self.panel
        self.assertTrue(H.wait_for(lambda: logical in p.profiles, 30))
        p._show_profile_details(logical)
        H.settle(0.1)
        pops = self._visible_details()
        self.assertEqual(len(pops), 1)
        return pops[0]

    def _reopened_details(self):
        """The popup the explorer opens again after a save dialog."""
        pops = self._visible_details()
        self.assertEqual(len(pops), 1)
        return pops[0]

    def test_29_statistics_popup_copy(self):
        from qgis.PyQt.QtGui import QGuiApplication
        from qgis.PyQt.QtTest import QTest
        pop = self._open_details(self.fl)
        clip = QGuiApplication.clipboard()
        # selectable summary: Ctrl+C copies the raw values of the selection
        from qgis.PyQt.QtCore import QItemSelectionModel
        view = pop.summary_view
        view.setFocus()
        flags = QItemSelectionModel.SelectionFlag
        view.selectionModel().select(view.model().index(0, 0),
                                     flags.ClearAndSelect | flags.Rows)
        QTest.keyClick(view, Qt.Key.Key_C, Qt.KeyboardModifier.ControlModifier)
        self.assertEqual(clip.text(), f"Records in scope\t{self.panel.profiles[self.fl].count}")
        self.assertTrue(pop.copied_label.isVisible())
        # distribution: selected cells, shares as fractions
        dist = pop.dist_view
        dist.setFocus()
        dist.selectAll()
        QTest.keyClick(dist, Qt.Key.Key_C, Qt.KeyboardModifier.ControlModifier)
        lines = clip.text().splitlines()
        self.assertEqual(len(lines), dist.topLevelItemCount())
        label, count, share = lines[0].split("\t")
        self.assertTrue(label.startswith("["))
        self.assertAlmostEqual(float(share), int(count) / self.panel.profiles[self.fl].valid)
        # the Copy menu
        acts = {a.text(): a for a in pop.copy_menu.actions()}
        self.assertEqual(list(acts), ["Summary Statistics", "Distribution Table", "Everything"])
        acts["Summary Statistics"].trigger()
        self.assertTrue(clip.text().startswith("Statistic\tValue\tShare\nRecords in scope"))
        self.assertIn("\nNon-NULL\t", clip.text())
        acts["Distribution Table"].trigger()
        self.assertTrue(clip.text().startswith("Bin\tRecords\tShare\n"))
        acts["Everything"].trigger()
        text = clip.text()
        for part in (f"Layer\t{self.layer.name()}", "Field\tSCORE_MAX (Decimal)",
                     "Active filters\tNone", "Statistic\tValue\tShare", "Bin\tRecords\tShare"):
            self.assertIn(part, text)
        pop.close()
        # a true/false column has no distribution table
        bool_idx = self.layer.fields().indexOf("IS_VACANT")
        pop = self._open_details(bool_idx)
        acts = {a.text(): a for a in pop.copy_menu.actions()}
        self.assertFalse(acts["Distribution Table"].isEnabled())
        self.assertIn("\ntrue\t", (acts["Summary Statistics"].trigger(),
                                      QGuiApplication.clipboard().text())[1])
        pop.close()

    def test_30_statistics_popup_export(self):
        import csv
        import tempfile
        from openpyxl import load_workbook
        from qgis.PyQt.QtWidgets import QFileDialog
        folder = tempfile.mkdtemp(prefix="gts_stats_")
        target = {}
        original = QFileDialog.getSaveFileName
        QFileDialog.getSaveFileName = staticmethod(
            lambda *a, **k: (os.path.join(folder, target["name"]), ""))
        p = self.panel
        p.header.cells[self.lu].set_filter(FL.ColumnFilter("LAND_USE", F.TEXT, "eq",
                                                           "RESIDENTIAL"))
        p.refresh()
        settle(p)
        try:
            pop = self._open_details(self.fl)
            acts = {a.text(): a for a in pop.export_menu.actions()}
            self.assertEqual(list(acts), ["Excel Workbook…", "Summary as CSV…",
                                          "Distribution as CSV…"])
            target["name"] = "stats.xlsx"
            acts["Excel Workbook…"].trigger()
            H.settle(0.3)
            wb = load_workbook(os.path.join(folder, "stats.xlsx"))
            self.assertEqual(wb.sheetnames, ["Summary", "Distribution"])
            summary = {r[0]: r[1:] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
            self.assertEqual(summary["Layer"][0], self.layer.name())
            self.assertIn("LAND_USE", summary["Active filters"][0])
            prof = p.profiles[self.fl]
            self.assertEqual(summary["Records in scope"][0], prof.count)
            self.assertAlmostEqual(summary["Non-NULL"][1], prof.non_null / prof.count)
            ws = wb["Distribution"]
            self.assertEqual([c.value for c in ws[1]], ["Bin", "Records", "Share"])
            share = ws.cell(2, 3)
            self.assertIsInstance(share.value, float)                # a number, not "<1%"
            self.assertEqual(share.number_format, "0.0%")
            # the popup opens again after the save dialog
            pop = self._reopened_details()
            target["name"] = "summary.csv"
            {a.text(): a for a in pop.export_menu.actions()}["Summary as CSV…"].trigger()
            H.settle(0.3)
            with open(os.path.join(folder, "summary.csv"), encoding="utf-8-sig") as fh:
                rows = list(csv.reader(fh))
            self.assertEqual(rows[0], ["Statistic", "Value", "Share"])
            self.assertEqual(rows[1], ["Records in scope", str(prof.count), ""])
            self.assertFalse(os.path.exists(os.path.join(folder, "summary_scope.txt")))
            pop = self._reopened_details()
            target["name"] = "dist.csv"
            {a.text(): a for a in pop.export_menu.actions()}["Distribution as CSV…"].trigger()
            H.settle(0.3)
            with open(os.path.join(folder, "dist.csv"), encoding="utf-8-sig") as fh:
                rows = list(csv.reader(fh))
            self.assertEqual(rows[0], ["Bin", "Records", "Share"])
            self.assertLessEqual(float(rows[1][2]), 1.0)
            self._reopened_details().close()
        finally:
            QFileDialog.getSaveFileName = original

    def test_21_replacement_keys(self):
        from qgis.gui import QgsGui
        from attribute_table_studio.core.compat import QAction
        mgr = QgsGui.shortcutsManager()
        native = QAction("Open Attribute Table", self.iface.mainWindow())
        native.setObjectName("mActionOpenTableTest")
        mgr.registerAction(native, "F6")
        mgr.setKeySequence(native, "F6")
        self.assertIs(mgr.actionForSequence(QKeySequence("F6")), native)
        settings.set_value("explorer/replaceTableKeys", True)
        try:
            self.plugin.apply_table_keys()
            self.assertTrue(native.shortcut().isEmpty())
            keys = {a.shortcut().toString() for a in self.plugin.table_keys}
            self.assertEqual(keys, {"F6", "Shift+F6", "Ctrl+F6"})
            self.plugin.table_keys[1].trigger()               # Shift+F6: selected
            panel = self.plugin.explorer_for(self.layer)
            self.assertEqual(panel.view_scope, "selected")
        finally:
            settings.set_value("explorer/replaceTableKeys", False)
            self.plugin.apply_table_keys()
        self.assertEqual(native.shortcut().toString(), "F6")   # given back
        self.assertEqual(self.plugin.table_keys, [])
        mgr.unregisterAction(native)


if __name__ == "__main__":
    H.run_tests(APP)
